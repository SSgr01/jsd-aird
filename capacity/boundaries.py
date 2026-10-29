"""Generate exact-size legal OOXML/PNG files; stream uploads only to the owned lab API."""
import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import struct
import zipfile
import zlib


def padded_ooxml(base, output, size):
    output = Path(output)
    padding = size - len(base) - 76 - 2 * len('capacity-padding.bin')
    if padding < 0:
        raise ValueError('Target smaller than document envelope')
    output.write_bytes(base)
    with zipfile.ZipFile(output, 'a', compression=zipfile.ZIP_STORED) as archive:
        with archive.open('capacity-padding.bin', 'w') as stream:
            remaining = padding
            while remaining:
                count = min(65536, remaining); stream.write(b'0' * count); remaining -= count
    if output.stat().st_size != size:
        raise ValueError('OOXML boundary fixture is not exact')


def padded_png(base, output, size):
    padding = size - len(base) - 12
    if padding < 0:
        raise ValueError('Target smaller than PNG envelope')
    with Path(output).open('wb') as stream:
        stream.write(base[:-12]); stream.write(struct.pack('>I', padding)); stream.write(b'npAD')
        crc = zlib.crc32(b'npAD'); remaining = padding
        while remaining:
            count = min(65536, remaining); chunk = b'0' * count
            stream.write(chunk); crc = zlib.crc32(chunk, crc); remaining -= count
        stream.write(struct.pack('>I', crc & 0xffffffff)); stream.write(base[-12:])
    if Path(output).stat().st_size != size:
        raise ValueError('PNG boundary fixture is not exact')


def run(run_id, password):
    import httpx
    from docx import Document
    from openpyxl import Workbook
    from PIL import Image
    output = Path('/results/boundaries'); output.mkdir(exist_ok=True)
    document = Document(); document.add_paragraph('隔离边界测试'); word = io.BytesIO(); document.save(word)
    book = Workbook(); book.active.append(['黏度']); book.active.append([100]); sheet = io.BytesIO(); book.save(sheet)
    image = io.BytesIO(); Image.new('RGB', (32,32), 'white').save(image, format='PNG')
    cases = [('KNOWLEDGE',512*1024*1024,'xlsx',sheet.getvalue()),
             ('PROJECT_DOCUMENT',200*1024*1024,'docx',word.getvalue()),
             ('EXPERIMENT_SOURCE',200*1024*1024,'docx',word.getvalue()),
             ('SPC_CHART',100*1024*1024,'png',image.getvalue()),
             ('SNAPSHOT',512*1024*1024,'xlsx',sheet.getvalue())]
    samples = []
    with httpx.Client(base_url='http://api:8080', timeout=httpx.Timeout(620, connect=3)) as client:
        csrf = client.get('/api/v1/auth/csrf').raise_for_status().json()['data']['token']
        client.headers['X-XSRF-TOKEN'] = csrf
        client.post('/api/v1/auth/login',json={'username':'capacity-0','password':password}).raise_for_status()
        for kind, maximum, extension, base in cases:
            for delta in (-1,0,1):
                size = maximum + delta
                if shutil.disk_usage(output).free < size + 10*1024**3:
                    raise RuntimeError('Disk safety floor prevents further boundary generation')
                path = output / f'{kind.lower()}-{delta+1}.{extension}'
                (padded_png if extension == 'png' else padded_ooxml)(base,path,size)
                digest = hashlib.sha256()
                with path.open('rb') as stream:
                    for chunk in iter(lambda:stream.read(65536),b''): digest.update(chunk)
                with path.open('rb') as stream:
                    response = client.post('/api/v1/files/staged', params={'kind':kind},files={'file':(path.name,stream,'application/octet-stream')})
                expected = 413 if delta > 0 else 200
                samples.append({'kind':kind,'bytes':size,'sha256':digest.hexdigest(),'status':response.status_code,'expected':expected,
                                'requestId':response.headers.get('x-request-id'),'passed':response.status_code==expected})
                (output/'results.json').write_text(json.dumps({'runId':run_id,'samples':samples},indent=2),encoding='utf-8')
                path.unlink() # Inputs can be regenerated; raw result/digest is retained.
                if response.status_code != expected:
                    raise RuntimeError(f'Boundary contract failed for {kind}, delta {delta}; see boundaries/results.json')


if __name__ == '__main__':
    run(os.environ['AIRD_CAPACITY_RUN_ID'], os.environ['CAPACITY_PASSWORD'])
