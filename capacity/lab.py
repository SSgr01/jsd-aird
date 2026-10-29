"""Only disposable Compose resources are accepted; no application .env is read."""
import hashlib
import io
import json
import os
from pathlib import Path
import re
import sys
import threading
import time
import uuid
from datetime import date

import httpx
import psycopg
from psycopg import sql

RUN = os.environ.get('AIRD_CAPACITY_RUN_ID', '')
if not re.fullmatch(r'aird-capacity-[a-z0-9-]{8,48}', RUN):
    raise SystemExit('A disposable capacity run ID is required')
if os.environ.get('PGHOST') != 'postgres' or os.environ.get('PGDATABASE') != 'aird_capacity':
    raise SystemExit('Only the isolated Compose database is allowed')
RESULTS = Path('/results')
ORG = uuid.UUID('00000000-0000-0000-0000-000000000001')


def connect():
    return psycopg.connect(application_name=RUN, connect_timeout=3)


def identifier(kind, number=0):
    return uuid.uuid5(uuid.NAMESPACE_URL, f'{RUN}/{kind}/{number}')


def save(name, value):
    RESULTS.mkdir(exist_ok=True)
    (RESULTS / name).write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str), encoding='utf-8')


def owned(connection):
    marker = connection.execute('SELECT run_id FROM public.capacity_run').fetchone()
    if not marker or marker[0] != RUN:
        raise RuntimeError('Database marker does not match this run')


def bootstrap():
    """Execute every original migration, recording source names/hashes, including duplicate numbers.

    This is an isolated test catalog, not a repair of any deployed Flyway history.
    Missing module foundations must be supplied as schema-only SQL in /assets/foundation.sql.
    """
    files = sorted(Path('/db/migration').glob('V*__*.sql'),
                   key=lambda p: (tuple(map(int, p.name.split('__')[0][1:].replace('_', '.').split('.'))), p.name))
    manifest = []
    with connect() as connection:
        if connection.execute("SELECT to_regclass('public.capacity_run')").fetchone()[0]:
            owned(connection)
            if not connection.execute('SELECT complete FROM public.capacity_run').fetchone()[0]:
                raise RuntimeError('Previous bootstrap is partial; use a fresh owned run instead of treating it as complete')
            return
        if connection.execute("SELECT to_regclass('iam.app_user')").fetchone()[0]:
            raise RuntimeError('Refusing to bootstrap a nonempty unmarked database')
        connection.execute('CREATE TABLE public.capacity_run(run_id text PRIMARY KEY,complete boolean NOT NULL DEFAULT false,created_at timestamptz DEFAULT now())')
        connection.execute('INSERT INTO public.capacity_run(run_id) VALUES(%s)', (RUN,))
        connection.commit()
        for path in files:
            text = path.read_text(encoding='utf-8-sig')
            connection.execute(text, prepare=False)
            connection.commit()
            manifest.append({'file': path.name, 'sha256': hashlib.sha256(text.encode()).hexdigest()})
            if path.name.startswith('V52__'):
                connection.execute(Path('/db/foundation/mdm_rnd_foundation.sql').read_text(encoding='utf-8-sig'), prepare=False)
                extra = Path('/assets/foundation.sql')
                if extra.exists():
                    raw = extra.read_text(encoding='utf-8-sig')
                    connection.execute(raw, prepare=False)
                    manifest.append({'file': 'operator-schema-only-foundation', 'sha256': hashlib.sha256(raw.encode()).hexdigest()})
                connection.commit()
        # The business API uses a non-superuser, with the same connection protocol.
        connection.execute(sql.SQL('CREATE ROLE capacity_app LOGIN PASSWORD {}').format(sql.Literal(os.environ['PGPASSWORD'])))
        schemas = ['core', 'iam', 'mdm', 'tpl', 'rnd', 'quality', 'spc', 'mfg', 'kb', 'ai', 'ops', 'data', 'public']
        for schema in schemas:
            connection.execute(sql.SQL('GRANT USAGE ON SCHEMA {} TO capacity_app').format(sql.Identifier(schema)))
            connection.execute(sql.SQL('GRANT ALL ON ALL TABLES IN SCHEMA {} TO capacity_app').format(sql.Identifier(schema)))
            connection.execute(sql.SQL('GRANT ALL ON ALL SEQUENCES IN SCHEMA {} TO capacity_app').format(sql.Identifier(schema)))
        connection.execute('UPDATE public.capacity_run SET complete=true WHERE run_id=%s',(RUN,))
        connection.commit()
    duplicates = sorted({p.name.split('__')[0] for p in files if sum(q.name.split('__')[0] == p.name.split('__')[0] for q in files) > 1})
    save('migration-manifest.json', {'runId': RUN, 'catalog': 'isolated SQL catalog; deployed Flyway history unchanged',
                                     'duplicateOriginalVersions': duplicates, 'scripts': manifest})


def insert(connection, table, values):
    connection.execute(sql.SQL('INSERT INTO {} ({}) VALUES ({})').format(
        sql.Identifier(*table.split('.')), sql.SQL(',').join(map(sql.Identifier, values)),
        sql.SQL(',').join(sql.Placeholder() for _ in values)), list(values.values()))


def files(connection):
    from minio import Minio
    from openpyxl import Workbook
    from docx import Document
    from PIL import Image
    client = Minio('minio:9000', access_key='capacity', secret_key=os.environ['CAPACITY_STORAGE_PASSWORD'], secure=False)
    if not client.bucket_exists(RUN):
        client.make_bucket(RUN)
    book = Workbook(); book.active.title='隔离'; book.active.append(['黏度']); book.active.append([100])
    workbook = io.BytesIO(); book.save(workbook)
    word = Document(); word.add_paragraph('隔离容量测试'); word.add_table(rows=2, cols=2).cell(0, 0).text = '样本'
    document = io.BytesIO(); word.save(document)
    image = io.BytesIO(); Image.new('RGB', (32, 32), 'white').save(image, format='PNG')
    artifacts = {'xlsx': (workbook.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                 'docx': (document.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
                 'png': (image.getvalue(), 'image/png'),
                 'json': (json.dumps({'id':'capacity-workbook','sheetOrder':['sheet1'],'sheets':{'sheet1':{'id':'sheet1','name':'隔离','rowCount':10,'columnCount':5,'cellData':{}}}},ensure_ascii=False,separators=(',',':')).encode(), 'application/json')}
    output = {}
    for extension, (data, mime) in artifacts.items():
        ids = []
        for n in range(400):
            file_id = identifier('file-' + extension, n); key = f'{RUN}/{file_id}/fixture.{extension}'
            client.put_object(RUN, key, io.BytesIO(data), len(data), content_type=mime)
            insert(connection, 'ops.file_object', dict(id=file_id, organization_id=ORG, bucket=RUN, object_key=key,
                   original_name=f'fixture-{n}.{extension}', content_type=mime, size_bytes=len(data),
                   sha256=hashlib.sha256(data).hexdigest(), status='ACTIVE', created_by=identifier('user')))
            ids.append(file_id)
        output[extension] = ids
    return output, artifacts


def seed():
    from argon2 import PasswordHasher
    with connect() as connection:
        owned(connection)
        required = ['mfg.inventory_product', 'mfg.inventory_stock_balance', 'mfg.inventory_transaction',
                    'mfg.inventory_policy', 'mfg.inventory_sample_dispatch', 'mfg.inventory_shipment',
                    'mfg.inventory_stock_import', 'quality.data_category', 'quality.data_record',
                    'quality.data_record_version', 'quality.upload_file']
        missing = [name for name in required if not connection.execute('SELECT to_regclass(%s)', (name,)).fetchone()[0]]
        save('schema-preflight.json', {'missingTables': missing, 'complete': not missing})
        if missing:
            raise RuntimeError('Full module fixtures blocked by missing base DDL; see schema-preflight.json. Supply schema-only /assets/foundation.sql and use a fresh run.')
        if connection.execute('SELECT 1 FROM iam.app_user WHERE username=%s', ('capacity-0',)).fetchone():
            raise RuntimeError('Fixtures already exist; use the existing manifest, or create a fresh run')
        role = identifier('role')
        insert(connection, 'iam.role', dict(id=role, organization_id=ORG, code='CAPACITY_TEST', name='隔离压测角色'))
        # Permission definitions are seeded by the real application's IAM bootstrap on startup.
        # Bindings are added after readiness by grant(), before any login or load.
        password = PasswordHasher(time_cost=2, memory_cost=16384, parallelism=1).hash(os.environ['CAPACITY_PASSWORD'])
        for n in range(500):
            insert(connection, 'iam.app_user', dict(id=identifier('user', n), organization_id=ORG, username=f'capacity-{n}',
                   display_name=f'隔离用户 {n}', role_id=role, password_hash=password, status='ACTIVE' if n<400 else 'DISABLED'))
        sources, artifacts = files(connection)
        empty_hash = hashlib.sha256(b'{}').hexdigest()
        manifest = {'runId': RUN, 'projects': [], 'experiments': [], 'templates': [], 'data': [],
                    'knowledge': [], 'research': [], 'production': [], 'inventory': [], 'quality': [], 'spectrum': []}
        for n in range(400):
            actor = identifier('user'); project = identifier('project', n); template = identifier('template', n); version = identifier('template-version', n)
            insert(connection, 'mdm.business_partner', dict(id=identifier('partner', n), partner_code=f'CAP-P-{n}',
                   name=f'隔离客户{n}', normalized_name=f'隔离客户{n}', status=['ACTIVE','INACTIVE'][n % 2], created_by='capacity-0', updated_by='capacity-0'))
            insert(connection, 'mdm.project', dict(id=project, project_code=f'CAP-{n}', name=f'隔离项目{n}',
                   status=['NOT_STARTED','IN_PROGRESS','COMPLETED'][n % 3], created_by='capacity-0', updated_by='capacity-0'))
            insert(connection, 'mdm.material', dict(id=identifier('material', n), code=f'CAP-M-{n}', name=f'隔离材料{n}',
                   category='隔离', source_category='MANUAL', source_module='CAPACITY', status='DRAFT', created_by='capacity-0', updated_by='capacity-0'))
            insert(connection, 'tpl.template', dict(id=template, organization_id=ORG, template_code=f'CAP-T-{n}', name=f'隔离模板{n}', format='XLSX', created_by=actor))
            insert(connection, 'tpl.template_version', dict(id=version, template_id=template, version_no=1, status='DRAFT', schema_jsonb='{"type":"object","properties":{}}',
                   editor_snapshot_file_id=sources['json'][n],editor_snapshot_hash=hashlib.sha256(artifacts['json'][0]).hexdigest(),
                   snapshot_kind='UNIVER_WORKBOOK', editor_app_version='capacity-fixture', plugin_manifest_hash=empty_hash,
                   schema_hash=empty_hash, mapping_hash=empty_hash, data_hash=empty_hash, workspace_hash=empty_hash, created_by=actor))
            experiment = identifier('experiment', n); exp_version = identifier('experiment-version', n)
            state = ['DRAFT','PENDING_REVIEW','COMPLETED'][n % 3]
            insert(connection, 'rnd.experiment', dict(id=experiment, organization_id=ORG, experiment_no=f'CAP-E-{n}',
                   title=f'隔离实验{n}', source_type='MANUAL', status=state, owner_name='capacity-0', created_by=actor, updated_by=actor))
            insert(connection, 'rnd.experiment_version', dict(id=exp_version, organization_id=ORG, experiment_id=experiment, version_no=1, status=state, created_by=actor))
            connection.execute('UPDATE rnd.experiment SET current_version_id=%s WHERE id=%s', (exp_version, experiment))
            research = identifier('research', n); research_version = identifier('research-version', n)
            insert(connection, 'rnd.research_test_record', dict(id=research, organization_id=ORG, record_type='REPORT', business_no=f'CAP-R-{n}',
                   name=f'隔离报告{n}', document_format='WORD', source_type='BLANK', status='DRAFT', created_by=actor, updated_by=actor))
            insert(connection, 'rnd.research_test_version', dict(id=research_version, organization_id=ORG, record_id=research, version_no=1, status='DRAFT', created_by=actor))
            connection.execute('UPDATE rnd.research_test_record SET current_version_id=%s WHERE id=%s', (research_version, research))
            job = identifier('data', n)
            data_template=identifier('data-template',n);data_version=identifier('data-template-version',n)
            insert(connection,'tpl.template',dict(id=data_template,organization_id=ORG,template_code=f'CAP-DT-{n}',name=f'隔离数据模板{n}',format='XLSX',created_by=actor))
            insert(connection,'tpl.template_version',dict(id=data_version,template_id=data_template,version_no=1,status='DRAFT',
                   schema_jsonb='{"type":"object","properties":{"viscosity":{"type":"number","dataType":"NUMBER","fieldCode":"VISCOSITY","dataPath":"/viscosity"}}}',
                   snapshot_kind='UNIVER_WORKBOOK',editor_app_version='capacity-fixture',plugin_manifest_hash=empty_hash,schema_hash=empty_hash,mapping_hash=empty_hash,data_hash=empty_hash,workspace_hash=empty_hash,created_by=actor))
            insert(connection, 'data.import_job', dict(id=job, organization_id=ORG, source_file_id=sources['xlsx'][n], source_sha256=hashlib.sha256(artifacts['xlsx'][0]).hexdigest(),
                   source_file_name=f'fixture-{n}.xlsx', source_format='XLSX', template_version_id=data_version,
                   status=['WAITING_MAPPING','WAITING_CONFIRM','FAILED'][n % 3], created_by=actor))
            sheet=identifier('data-sheet',n)
            insert(connection,'data.import_sheet',dict(id=sheet,import_job_id=job,sheet_id='sheet1',sheet_name='隔离',sheet_order=0,selected=True,
                   header_rows_jsonb='[1]',data_start_row=2,data_end_row=2,confirmation_status='CONFIRMED'))
            insert(connection,'data.import_mapping',dict(id=identifier('data-mapping',n),import_job_id=job,import_sheet_id=sheet,source_column='A',source_header='黏度',
                   field_code='VISCOSITY',field_name='黏度',action='MAP',value_type='NUMBER',mapping_jsonb='{"bindingId":"viscosity","dataPath":"/viscosity"}',status='CONFIRMED'))
            insert(connection,'data.staging_row',dict(id=identifier('data-row',n),import_job_id=job,import_sheet_id=sheet,source_row_number=2,raw_values_jsonb='{"A":"100"}',status='STAGED'))
            document = identifier('knowledge', n)
            insert(connection, 'kb.document', dict(id=document, organization_id=ORG, title=f'隔离知识{n}', document_type='DOCX',
                   status=['READY','FAILED','QUEUED'][n % 3], scan_status='SAFE', created_by=actor))
            insert(connection, 'kb.document_version', dict(id=identifier('knowledge-version', n), document_id=document, version_no=1,
                   file_object_id=sources['docx'][n], original_name=f'fixture-{n}.docx', content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                   size_bytes=len(artifacts['docx'][0]), sha256=hashlib.sha256(artifacts['docx'][0]).hexdigest(), status=['READY','FAILED','QUEUED'][n % 3]))
            chart = identifier('spectrum', n)
            category=identifier('spectrum-category',n)
            insert(connection,'spc.chart_category',dict(id=category,organization_id=ORG,code=f'CAP-S-{n}',name=f'隔离图谱类别{n}',description='隔离测试',created_by=actor))
            insert(connection, 'spc.chart_asset', dict(id=chart, organization_id=ORG, file_object_id=sources['png'][n],
                   category_id=category,title=f'隔离图谱{n}', original_name=f'fixture-{n}.png', content_type='image/png',
                   size_bytes=len(artifacts['png'][0]),sha256=hashlib.sha256(artifacts['png'][0]).hexdigest(),created_by=actor))
            product = identifier('inventory', n)
            insert(connection, 'mfg.inventory_product', dict(id=product, organization_id=ORG, code=f'CAP-I-{n}', name=f'隔离库存{n}',
                   category='隔离', inventory_scope='RND', status='ACTIVE', created_by=actor))
            insert(connection, 'mfg.inventory_stock_balance', dict(id=identifier('balance', n), organization_id=ORG, product_id=product,
                   inventory_scope='RND', quantity_kg=100000))
            quality = identifier('quality', n)
            insert(connection, 'quality.data_record', dict(id=quality, organization_id=ORG, business_type='record', business_no=f'CAP-Q-{n}',
                   data_jsonb='{}', created_by=actor, updated_by=actor))
            order = identifier('production', n)
            insert(connection, 'mfg.production_order', dict(id=order, organization_id=ORG, order_no=f'CAP-O-{n}', status='DRAFT', template_version_id=version,
                   instance_schema_jsonb='{}', instance_mapping_jsonb='[]', snapshot_kind='UNIVER_WORKBOOK', editor_app_version='capacity-fixture',
                   plugin_manifest_hash=empty_hash, schema_hash=empty_hash, mapping_hash=empty_hash, data_hash=empty_hash, workspace_hash=empty_hash, created_by=actor))
            for key, value in [('projects', project), ('templates', version), ('experiments', experiment), ('data', job), ('knowledge', document),
                               ('research', research), ('inventory', product), ('quality', quality), ('production', order), ('spectrum', chart)]:
                manifest[key].append(str(value))
        manifest['reviewers']=[str(identifier('user',n)) for n in range(400)]
        connection.commit()
    save('fixtures.json', manifest)


def grant():
    with connect() as connection:
        owned(connection)
        connection.execute("""INSERT INTO iam.role_permission_binding(id,organization_id,role_id,permission_code,effect,scope_type)
            SELECT gen_random_uuid(),%s,%s,code,'ALLOW','ALL' FROM iam.permission_definition
            ON CONFLICT(organization_id,role_id,permission_code) DO NOTHING""", (ORG, identifier('role')))


def tier(module, count):
    if module not in ('projects', 'experiments', 'data') or count not in (10000, 100000, 1000000):
        raise ValueError('Tier must select exactly one allowlisted table and 10k/100k/1m rows')
    with connect() as connection:
        owned(connection)
        # Rotate the large table. Only rows bearing this synthetic marker are deleted.
        for table, column in [('mdm.project', 'project_code'), ('rnd.experiment', 'experiment_no'), ('data.data_record', 'record_key')]:
            connection.execute(sql.SQL("DELETE FROM {} WHERE {} LIKE 'CAP-LARGE-%'").format(sql.Identifier(*table.split('.')), sql.Identifier(column)))
        if module == 'projects':
            connection.execute("""INSERT INTO mdm.project(id,project_code,name,status,created_by,updated_by)
                SELECT gen_random_uuid(),'CAP-LARGE-'||n,'隔离大表项目'||n,'NOT_STARTED','capacity-0','capacity-0' FROM generate_series(1,%s) n""", (count,))
        elif module == 'experiments':
            connection.execute("""INSERT INTO rnd.experiment(id,organization_id,experiment_no,title,source_type,status,created_by,updated_by)
                SELECT gen_random_uuid(),%s,'CAP-LARGE-'||n,'隔离大表实验'||n,'MANUAL','DRAFT',%s,%s FROM generate_series(1,%s) n""", (ORG, identifier('user'), identifier('user'), count))
        else:
            connection.execute("""INSERT INTO data.data_record(id,organization_id,import_job_id,record_key,record_index,raw_data_jsonb,normalized_data_jsonb,corrected_data_jsonb,effective_data_jsonb,quality_status,synthetic_key)
                SELECT gen_random_uuid(),%s,%s,'CAP-LARGE-'||n,n,'{}','{}','{}','{}','VALID',false FROM generate_series(1,%s) n""", (ORG, identifier('data'), count))
        connection.execute('ANALYZE')
    save('dataset-tier.json', {'module': module, 'rows': count})


def database_metrics():
    with connect() as connection:
        owned(connection)
        activity = connection.execute("""SELECT count(*),count(*) FILTER(WHERE wait_event_type='Lock'),
            coalesce(max(extract(epoch FROM clock_timestamp()-xact_start)) FILTER(WHERE state<>'idle'),0)
            ,count(*) FILTER(WHERE state<>'idle')
            FROM pg_stat_activity WHERE datname=current_database() AND pid<>pg_backend_pid()""").fetchone()
        jobs = connection.execute("SELECT status,count(*),coalesce(max(extract(epoch FROM now()-created_at)),0) FROM ops.async_job GROUP BY status").fetchall()
        return {'connectionsTotal': activity[0], 'connectionsActive':activity[3], 'lockWaiters': activity[1], 'oldestTransactionSeconds': float(activity[2]), 'jobs': jobs}


def fault(kind):
    if kind in ('slow-sql','connection-pool'):
        with connect() as connection:
            owned(connection)
            connection.execute("CREATE FUNCTION public.capacity_slow_read() RETURNS boolean LANGUAGE plpgsql AS $$ BEGIN PERFORM pg_sleep(3); RETURN true; END $$")
            connection.execute('ALTER TABLE mdm.project ENABLE ROW LEVEL SECURITY')
            connection.execute('CREATE POLICY capacity_slow_read ON mdm.project FOR SELECT TO capacity_app USING(public.capacity_slow_read())')
        try:
            if kind == 'slow-sql':
                time.sleep(20)
            else:
                from concurrent.futures import ThreadPoolExecutor
                with httpx.Client(base_url='http://api:8080',timeout=15) as client:
                    token=client.get('/api/v1/auth/csrf').raise_for_status().json()['data']['token']
                    client.headers['X-XSRF-TOKEN']=token
                    client.post('/api/v1/auth/login',json={'username':'capacity-0','password':os.environ['CAPACITY_PASSWORD']}).raise_for_status()
                    barrier=threading.Barrier(32)
                    def read(_):
                        barrier.wait(timeout=5)
                        response=client.get('/api/v1/projects?size=20')
                        return {'status':response.status_code,'code':response.json().get('code'),'requestId':response.headers.get('x-request-id')}
                    with ThreadPoolExecutor(max_workers=32) as pool: outcomes=list(pool.map(read,range(32)))
                    save('fault-connection-pool.json',{'responses':outcomes})
                    if not any(item['status']==503 and item['code']=='DATABASE_BUSY' for item in outcomes):
                        raise RuntimeError('Connection-pool saturation did not produce DATABASE_BUSY; inspect outcomes')
        finally:
            with connect() as connection:
                owned(connection)
                connection.execute('DROP POLICY IF EXISTS capacity_slow_read ON mdm.project')
                connection.execute('ALTER TABLE mdm.project DISABLE ROW LEVEL SECURITY')
                connection.execute('DROP FUNCTION IF EXISTS public.capacity_slow_read()')
    elif kind in ('storage','quota'):
        with httpx.Client(base_url='http://api:8080',timeout=40) as client:
            token=client.get('/api/v1/auth/csrf').raise_for_status().json()['data']['token'];client.headers['X-XSRF-TOKEN']=token
            client.post('/api/v1/auth/login',json={'username':'capacity-0','password':os.environ['CAPACITY_PASSWORD']}).raise_for_status()
            if kind=='storage': response=client.get(f'/api/v1/files/{identifier("file-xlsx")}/content')
            else: response=client.post('/api/v1/files/staged',files={'file':('sample.txt',b'isolated quota check','text/plain')})
            save('fault-'+kind+'.json',{'status':response.status_code,'retryAfter':response.headers.get('retry-after'),'requestId':response.headers.get('x-request-id')})
            if kind=='quota' and (response.status_code!=503 or response.json().get('code')!='SERVER_BUSY'):
                raise RuntimeError('Exhausted temporary quota did not reject admission')
            if kind=='storage' and response.status_code not in (503,504): raise RuntimeError('Storage fault did not return a bounded resource failure')
    elif kind == 'long-lock':
        with connect() as connection:
            owned(connection)
            connection.execute('SELECT id FROM mdm.project WHERE id=%s FOR UPDATE', (identifier('project'),))
            connection.execute('SELECT pg_sleep(20)')
    elif kind == 'deadlock':
        barrier = threading.Barrier(2); errors = []
        def competitor(first, second):
            try:
                with connect() as connection:
                    owned(connection)
                    connection.execute('UPDATE mdm.project SET version=version+1 WHERE id=%s', (identifier('project', first),))
                    barrier.wait(timeout=5)
                    connection.execute('UPDATE mdm.project SET version=version+1 WHERE id=%s', (identifier('project', second),))
            except psycopg.Error as error:
                errors.append(error.sqlstate)
        threads = [threading.Thread(target=competitor, args=pair) for pair in [(0,1),(1,0)]]
        for thread in threads: thread.start()
        for thread in threads: thread.join(8)
        if errors != ['40P01'] or any(t.is_alive() for t in threads): raise RuntimeError('Deadlock did not resolve as expected')
        save('fault-deadlock.json', {'sqlstates': errors})
    elif kind == 'queue':
        with connect() as connection:
            owned(connection)
            for n in range(20):
                connection.execute("INSERT INTO ops.async_job(id,organization_id,job_type,status,idempotency_key,next_attempt_at) VALUES(%s,%s,'CAPACITY_NO_HANDLER','READY',%s,now()+interval '1 hour')", (identifier('queue', n), ORG, f'{RUN}-queue-{n}'))
        rejected = False
        try:
            with connect() as connection:
                connection.execute("INSERT INTO ops.async_job(id,organization_id,job_type,status) VALUES(%s,%s,'CAPACITY_NO_HANDLER','READY')", (uuid.uuid4(), ORG))
        except psycopg.Error as error:
            rejected = error.sqlstate == '55P03'
        finally:
            with connect() as connection:
                connection.execute("DELETE FROM ops.async_job WHERE idempotency_key LIKE %s", (RUN + '-queue-%',))
        if not rejected: raise RuntimeError('Queue overflow was accepted')
    else:
        raise ValueError('Unknown fault')


if __name__ == '__main__':
    command = sys.argv[1]
    if command == 'bootstrap': bootstrap()
    elif command == 'seed': seed()
    elif command == 'grant': grant()
    elif command == 'tier': tier(sys.argv[2], int(sys.argv[3]))
    elif command == 'metrics':
        value=database_metrics()
        if len(sys.argv)>2:
            name=sys.argv[2]
            if name not in ('baseline','recovery'):raise ValueError('Unknown metric checkpoint')
            save(name+'.json',value)
        else:print(json.dumps(value, default=str))
    elif command == 'fault': fault(sys.argv[2])
    elif command == 'boundaries':
        from boundaries import run
        run(RUN,os.environ['CAPACITY_PASSWORD'])
    elif command == 'compress':
        import gzip
        import shutil
        name=sys.argv[2]
        if not re.fullmatch(r'[a-z0-9-]+',name): raise ValueError('Invalid stage name')
        source=RESULTS/(name+'.ndjson')
        with source.open('rb') as input, gzip.open(str(source)+'.gz','wb',compresslevel=5) as output: shutil.copyfileobj(input,output,65536)
        source.unlink()
    else: raise SystemExit('Unknown isolated lab command')
