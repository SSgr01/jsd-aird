from __future__ import annotations

import hashlib
import importlib.util
import json
import math
import random
import shutil
import zipfile
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

BASE = Path('/mnt/data/74137de1-ac32-433f-8bbf-98e589d34cd6.xlsx')
GEN = Path('/mnt/data/UVPU_APPLICATION_FORMULATION_Synthetic_TrainingSnapshot_V1/generate_uvpu_synthetic_snapshot.py')
OUT_XLSX = Path('/mnt/data/UVPU_APPLICATION_FORMULATION_Synthetic_400_Shared_Process_Original_Report_V3.xlsx')
SNAP_DIR = Path('/mnt/data/UVPU_APPLICATION_FORMULATION_Synthetic_TrainingSnapshot_V2_Shared_Process')
OUT_ZIP = Path('/mnt/data/UVPU_APPLICATION_FORMULATION_Synthetic_Shared_Process_V3_Package.zip')
SEED = 20260903
SNAPSHOT_ID = 'SYNTH-UVPU-APP-SHARED-PROCESS-20260903-V2'
GEN_VERSION = '2.0.0-shared-process'

spec = importlib.util.spec_from_file_location('uvpu_gen', GEN)
mod = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(mod)
base_measurements, base_source_map = mod.generate_rows(SEED)
assert len(base_measurements) == 400

# Reorder to make each report sheet contain multiple formulation/resin variants,
# while preserving formula-lineage ids for grouped validation.
# Original generator layout = 40 lineages x 10 replicates. Interleave lineages by replicate.
order = []
for rep in range(mod.REPLICATES_PER_LINEAGE):
    for lin in range(mod.N_LINEAGES):
        order.append(lin * mod.REPLICATES_PER_LINEAGE + rep)
assert len(order) == 400 and len(set(order)) == 400

records = []
for idx in order:
    m = dict(base_measurements[idx])
    s = dict(base_source_map[idx])
    records.append((m, s))

RESIN_DISPLAY = {
    'SJ-230': 'SJ-230',
    'SJ-231': 'SJ-231',
    'SJ-232': 'SJ-232',
    'SJ-230-WASHED': 'SJ-230水洗后',
    'SJ-231-WASHED': 'SJ-231水洗后',
    'SJ-232-WASHED': 'SJ-232水洗后',
}
RESIN_PROP = {
    'SJ-230': {'appearance':'微黄透明液体','solids':0.763,'dry':'粘手不湿手','viscosity':250,'water':0.00324,'mw':3173,'coating':'无色透明液体','film':'无色透明'},
    'SJ-231': {'appearance':'黄色透明液体','solids':0.723,'dry':'较粘手，湿手','viscosity':120,'water':0.00445,'mw':1674,'coating':'淡黄色透明液体','film':'淡黄色透明'},
    'SJ-232': {'appearance':'黄棕色微浑半透液体','solids':0.759,'dry':'粘手不湿手','viscosity':300,'water':0.00312,'mw':3235,'coating':'黄色透明液体','film':'淡黄色透明'},
    'SJ-230-WASHED': {'appearance':'白色微浑液体','solids':0.754,'dry':'粘手不湿手','viscosity':260,'water':0.02484,'mw':3050,'coating':'淡黄色透明液体','film':'无色透明'},
    'SJ-231-WASHED': {'appearance':'黄色微浑液体','solids':0.760,'dry':'较粘手，湿手','viscosity':150,'water':0.00136,'mw':1650,'coating':'无色透明液体','film':'无色透明'},
    'SJ-232-WASHED': {'appearance':'黄棕色微浑液体','solids':0.576,'dry':'粘手不湿手','viscosity':200,'water':0.09332,'mw':3100,'coating':'黄色微浑液体','film':'淡黄色微雾'},
}
FORMULA_ROWS = {
    'X_FORMULA__SJ_230_PCT':17,
    'X_FORMULA__SJ_231_PCT':18,
    'X_FORMULA__SJ_232_PCT':19,
    'X_FORMULA__SJ_230_WASHED_PCT':20,
    'X_FORMULA__SJ_231_WASHED_PCT':21,
    'X_FORMULA__SJ_232_WASHED_PCT':22,
    'X_FORMULA__DSP_3315_PCT':23,
    'X_FORMULA__SS059_PC_1_1_PCT':24,
    'X_FORMULA__S_48_PCT':25,
}
FORMULA_LABELS = {
    17:'SJ-230',18:'SJ-231',19:'SJ-232',20:'SJ-230水洗后',21:'SJ-231水洗后',22:'SJ-232水洗后',
    23:'DSP-3315',24:'SS059/碳酸丙烯酯=1/1',25:'S-48'
}
HARDNESS = {1.0:'H',2.0:'2H',3.0:'3H',4.0:'4H'}

# Discrete, realistic development-only public process conditions.
INTENSITY_LEVELS = [145.0, 155.0, 165.0, 175.0, 185.0]
ENERGY_LEVELS = [650.0, 700.0, 750.0, 800.0, 850.0, 900.0, 950.0]
TEMP_LEVELS = [23.0, 24.5, 26.0, 27.5, 29.0]
RH_LEVELS = [50.0, 55.0, 60.0, 65.0, 70.0, 75.0]
SOLIDS_LEVELS = [38.0, 39.0, 40.0, 41.0, 42.0]


def missing(v):
    return isinstance(v, float) and math.isnan(v)


def clamp(v, lo, hi):
    return lo if v < lo else hi if v > hi else v


def noise(exp_id: str, key: str, sd: float):
    h = hashlib.sha256(f'{SEED}:{GEN_VERSION}:{exp_id}:{key}'.encode('utf-8')).digest()
    rr = random.Random(int.from_bytes(h[:8], 'big'))
    return rr.gauss(0.0, sd)


def report_process(report_idx: int):
    rr = random.Random(SEED + report_idx * 7919)
    # Deterministic but shuffled across the designed level sets.
    return {
        'intensity': rr.choice(INTENSITY_LEVELS),
        'energy': rr.choice(ENERGY_LEVELS),
        'temp': rr.choice(TEMP_LEVELS),
        'rh': rr.choice(RH_LEVELS),
        'solids': rr.choice(SOLIDS_LEVELS),
    }


def recompute_targets(m: dict, s: dict, proc: dict):
    resin = s['main_resin_code']
    anchor = mod.RESIN_ANCHORS[resin]
    main_col = mod.FORMULA_COLS[resin]
    main_pct = float(m[main_col])
    dsp = float(m['X_FORMULA__DSP_3315_PCT'])
    ss = float(m['X_FORMULA__SS059_PC_1_1_PCT'])
    thickness = float(m['X_PROCESS__FILM_THICKNESS_UM'])
    exp = m['experiment_version_id']
    energy = proc['energy']
    intensity = proc['intensity']
    temp = proc['temp']
    rh = proc['rh']
    delta_main = main_pct - anchor['ref_pct']

    warp = (
        anchor['warp'] + 0.035*delta_main + 0.16*(dsp-1.60) + 0.10*(ss-1.23)
        - 0.0030*(energy-800.0) - 0.0035*(intensity-165.0)
        + 0.09*(thickness-9.5) + 0.012*(rh-67.0) + 0.025*(temp-26.0)
        + 0.0008*(energy-800.0)*(dsp-1.60) + noise(exp,'warp',0.22)
    )
    warp = clamp(warp,0.0,6.5)
    warp12 = (
        anchor['warp12'] + 0.62*(warp-anchor['warp']) + 0.018*(rh-67.0)
        + 0.04*(thickness-9.5) + noise(exp,'warp12',0.28)
    )
    warp12 = clamp(warp12,0.0,7.5)
    hardness_latent = (
        2.05 + anchor['hardness_shift'] + 0.0042*(energy-800.0) + 0.0040*(intensity-165.0)
        + 0.16*(dsp-1.60) + 0.025*delta_main - 0.012*(rh-67.0)
        + 0.035*(thickness-9.5) + noise(exp,'hardness',0.22)
    )
    hardness = float(int(clamp(round(hardness_latent),1,4)))
    abrasion500 = (
        anchor['abrasion500'] + 0.95*(energy-800.0) + 1.65*(intensity-165.0)
        + 52.0*(hardness-2.0) + 4.0*delta_main + 35.0*(dsp-1.60)
        - 1.7*(rh-67.0) + noise(exp,'abr500',35.0)
    )
    abrasion500 = clamp(abrasion500,5.0,800.0)
    abrasion1kg = (
        0.235*abrasion500 + 18.0*(hardness-2.0) + 0.30*(energy-800.0)
        + noise(exp,'abr1kg',16.0)
    )
    abrasion1kg = clamp(abrasion1kg,5.0,350.0)

    # Preserve the original missing/not-tested pattern, but recompute available values.
    if not missing(m['Y__WARPING_PET_INITIAL_CM']): m['Y__WARPING_PET_INITIAL_CM'] = round(warp,4)
    if not missing(m['Y__WARPING_PET_12H_CM']): m['Y__WARPING_PET_12H_CM'] = round(warp12,4)
    if not missing(m['Y__HARDNESS_PET_1KG_ORD']): m['Y__HARDNESS_PET_1KG_ORD'] = hardness
    if not missing(m['Y__STEEL_WOOL_500G_CYCLES']): m['Y__STEEL_WOOL_500G_CYCLES'] = round(abrasion500,4)
    if not missing(m['Y__STEEL_WOOL_1KG_CYCLES']): m['Y__STEEL_WOOL_1KG_CYCLES'] = round(abrasion1kg,4)

    # Public process is copied to every experiment row in the model snapshot.
    m['X_PROCESS__COATING_SOLIDS_PCT'] = proc['solids']
    m['X_PROCESS__UVA_INTENSITY_MW_CM2'] = proc['intensity']
    m['X_PROCESS__UV_ENERGY_MJ_CM2'] = proc['energy']
    m['X_ENV__TEMPERATURE_C'] = proc['temp']
    m['X_ENV__HUMIDITY_RH_PCT'] = proc['rh']


def resin_props(m, s):
    code = s['main_resin_code']
    base = RESIN_PROP[code]
    exp = m['experiment_version_id']
    return {
        'appearance': base['appearance'],
        'solids': clamp(base['solids'] + noise(exp,'resin_solids',0.009),0.50,0.85),
        'dry': base['dry'],
        'viscosity': max(60.0, base['viscosity']*(1+noise(exp,'viscosity_rel',0.10))),
        'water': max(0.0004, base['water']*(1+noise(exp,'water_rel',0.18))),
        'mw': max(900.0, base['mw']*(1+noise(exp,'mw_rel',0.045))),
        'coating': base['coating'],
        'film': base['film'],
    }


def fmt_warp(v):
    if missing(v): return '/（未测试）'
    if v < 0.08: return '平整，不翘'
    return f'反翘{float(v):.2f}cm'


def fmt_hardness(v):
    if missing(v): return '/（未测试）'
    label = HARDNESS[float(v)]
    nxt={'H':'2H','2H':'3H','3H':'4H','4H':'5H'}[label]
    return f'{label}，OK\n{nxt}，轻微痕'


def fmt_cycles(v):
    if missing(v): return '/（未测试）'
    n=int(round(float(v)))
    desc='无明显丝痕' if n>=450 else '少量浅丝痕' if n>=220 else '轻微丝痕' if n>=80 else '出现明显丝痕'
    return f'{n}次\n{desc}'

# Apply one explicit public process to every record in each report sheet.
chunks=[]
for i in range(0,400,6):
    chunks.append(records[i:i+6])
assert len(chunks)==67 and len(chunks[-1])==4
for report_idx, chunk in enumerate(chunks,1):
    proc=report_process(report_idx)
    sheet=f'APP-{report_idx:03d}'
    for col_idx,(m,s) in enumerate(chunk, start=4):
        recompute_targets(m,s,proc)
        s['source_file'] = OUT_XLSX.name
        s['source_sheet'] = sheet
        s['source_range'] = f'{get_column_letter(col_idx)}8:{get_column_letter(col_idx)}36'
        s['content_hash'] = mod.row_hash(m)
        s['generator_version'] = GEN_VERSION

# ---------- Create original-format workbook ----------
wb=load_workbook(BASE)
for name in list(wb.sheetnames):
    if name!='Sheet1': del wb[name]
tmpl=wb['Sheet1']
tmpl.title='__TEMPLATE'

for report_idx, chunk in enumerate(chunks,1):
    proc=report_process(report_idx)
    ws=wb.copy_worksheet(tmpl)
    ws.title=f'APP-{report_idx:03d}'
    ws.sheet_view.showGridLines=False
    ws['A1']='UV/PU应用配方性能模拟测试（SYNTHETIC）'
    ws['C2']=f'模拟实验员-{((report_idx-1)%4)+1:02d}'
    day=((report_idx-1)%28)+1; month=7+((report_idx-1)//28)%2
    ws['E2']=f'2026.{month}.{day}'
    ws['G2']=f"{proc['temp']:.1f}°C"
    ws['I2']=f"{proc['rh']:.0f}%RH"
    ws['B3']='UV/PU应用配方性能模拟验证'
    ws['B4']='在同一公共工艺条件下比较不同配方的应用性能；用于T03–T07开发、导入、建模与优化流程测试'
    ws['B5']='测试用素材：100μm厚的光学级PET膜（本页D:I实验共用）'
    ws['B6']='制膜施工方式：标格达40μm绕丝棒辊涂（本页D:I实验共用）'
    ws['B7']=(
        f"固化条件：辊涂后，汞灯UV固化，UVA光强：{proc['intensity']:.0f}mW/cm²，"
        f"UVA能量：{proc['energy']:.0f}mJ/cm²（本页D:I实验共用；建模时复制到每条实验X）"
    )
    ws.row_dimensions[7].height=40

    for rr,label in FORMULA_LABELS.items(): ws.cell(rr,3).value=label
    ws['C26']='合计'
    for c in range(4,10):
        for rr in list(range(8,16))+list(range(16,37)):
            if rr==26: continue
            ws.cell(rr,c).value=None
            ws.cell(rr,c).comment=None
        col=get_column_letter(c)
        ws.cell(26,c).value=f'=IF(COUNT({col}17:{col}25)=0,"",SUM({col}17:{col}25))'
        ws.cell(26,c).number_format='0.0000'

    for offset,(m,s) in enumerate(chunk):
        c=4+offset
        p=resin_props(m,s)
        exp_id=m['experiment_version_id'].replace('-V1','')
        resin=RESIN_DISPLAY[s['main_resin_code']]
        ws.cell(8,c).value=resin
        ws.cell(9,c).value=f'主树脂模拟批次 {exp_id[-4:]}'
        ws.cell(10,c).value=p['appearance']
        ws.cell(11,c).value=round(p['solids'],4); ws.cell(11,c).number_format='0.00%'
        ws.cell(12,c).value=p['dry']
        ws.cell(13,c).value=int(round(p['viscosity']))
        ws.cell(14,c).value=round(p['water'],5); ws.cell(14,c).number_format='0.000%'
        ws.cell(15,c).value=int(round(p['mw']))
        ws.cell(16,c).value=exp_id
        for field,rr in FORMULA_ROWS.items():
            ws.cell(rr,c).value=round(float(m[field]),4); ws.cell(rr,c).number_format='0.0000'
        ws.cell(27,c).value=proc['solids']/100.0; ws.cell(27,c).number_format='0.00%'
        ws.cell(28,c).value=p['coating']
        ws.cell(29,c).value='轻微粘手' if proc['energy']<700 else 'OK'
        film=p['film'] + ('，轻微雾影' if proc['rh']>=75 and '透明' in p['film'] else '')
        ws.cell(30,c).value=film
        ws.cell(31,c).value=round(float(m['X_PROCESS__FILM_THICKNESS_UM']),2); ws.cell(31,c).number_format='0.00'
        ws.cell(32,c).value=fmt_warp(m['Y__WARPING_PET_INITIAL_CM'])
        ws.cell(33,c).value=fmt_warp(m['Y__WARPING_PET_12H_CM'])
        ws.cell(34,c).value=fmt_hardness(m['Y__HARDNESS_PET_1KG_ORD'])
        ws.cell(35,c).value=fmt_cycles(m['Y__STEEL_WOOL_500G_CYCLES'])
        ws.cell(36,c).value=fmt_cycles(m['Y__STEEL_WOOL_1KG_CYCLES'])

    valid500=[(m['Y__STEEL_WOOL_500G_CYCLES'],m['experiment_version_id'].replace('-V1','')) for m,s in chunk if not missing(m['Y__STEEL_WOOL_500G_CYCLES'])]
    validwarp=[(m['Y__WARPING_PET_INITIAL_CM'],m['experiment_version_id'].replace('-V1','')) for m,s in chunk if not missing(m['Y__WARPING_PET_INITIAL_CM'])]
    parts=[f"公共工艺：UVA {proc['intensity']:.0f}mW/cm² / {proc['energy']:.0f}mJ/cm² / {proc['temp']:.1f}°C / {proc['rh']:.0f}%RH / 涂料固含{proc['solids']:.0f}%"]
    if valid500:
        best=max(valid500,key=lambda x:x[0]); parts.append(f'500g耐磨最高{int(round(best[0]))}次（{best[1]}）')
    if validwarp:
        best=min(validwarp,key=lambda x:x[0]); parts.append(f'即时翘曲最低{best[0]:.2f}cm（{best[1]}）')
    ws['A38']='结果/结论/小结：'+'；'.join(parts)+'。'
    ws['A38'].alignment=Alignment(horizontal='left',vertical='center',wrap_text=True); ws.row_dimensions[38].height=48
    ws['A39']='备注：本表全部为SYNTHETIC模拟数据，仅用于系统开发和模型流程测试；同一Sheet内D:I实验共享B5–B7及第2行温湿度公共工艺，进入建模层后必须把这些公共条件复制到每条实验X；“/（未测试）”为缺失，绝不等于0。'
    ws['A39'].font=Font(name=ws['A39'].font.name or '宋体',size=ws['A39'].font.sz or 11,bold=True,color='C00000')
    ws['A39'].alignment=Alignment(horizontal='left',vertical='center',wrap_text=True); ws.row_dimensions[39].height=68
    if len(chunk)<6:
        for c in range(4+len(chunk),10): ws.column_dimensions[get_column_letter(c)].hidden=True

# Remove template.
del wb['__TEMPLATE']
wb.active=0
try:
    wb.calculation.calcMode='auto'; wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True
except Exception:
    pass
wb.save(OUT_XLSX)

# ---------- Build aligned TrainingSnapshot V2 ----------
SNAP_DIR.mkdir(parents=True,exist_ok=True)
measurements=[m for m,s in records]
source_map=[s for m,s in records]

# Formula checks.
for m in measurements:
    total=sum(float(m[c]) for c in mod.FORMULA_COLS.values())+float(m['X_FORMULA__DSP_3315_PCT'])+float(m['X_FORMULA__SS059_PC_1_1_PCT'])+float(m['X_FORMULA__S_48_PCT'])
    assert abs(total-100.0)<0.02

meas_path=SNAP_DIR/'measurements.parquet'
src_path=SNAP_DIR/'source-map.parquet'
mod.write_minimal_parquet(meas_path,measurements,mod.MEASUREMENT_COLUMNS)
mod.write_minimal_parquet(src_path,source_map,mod.SOURCE_MAP_COLUMNS)
mod.write_csv(SNAP_DIR/'measurements_preview.csv',measurements,mod.MEASUREMENT_COLUMNS,limit=50)
mod.write_csv(SNAP_DIR/'source-map_preview.csv',source_map,mod.SOURCE_MAP_COLUMNS,limit=50)

# Custom manifest to document shared-process semantics.
target_stats={}
for col,definition in mod.TARGETS.items():
    vals=[r[col] for r in measurements]
    valid=[float(v) for v in vals if not missing(v)]
    target_stats[col]={
        'valid_count':len(valid),'missing_count':len(vals)-len(valid),'missing_rate':round((len(vals)-len(valid))/len(vals),4),
        'min':round(min(valid),4) if valid else None,'max':round(max(valid),4) if valid else None,
    }
feature_schema=[]
for col in mod.MEASUREMENT_COLUMNS:
    role='TARGET' if col.startswith('Y__') else 'IDENTITY' if col=='experiment_version_id' else 'FEATURE'
    sample=measurements[0][col]
    ftype='DOUBLE' if isinstance(sample,(int,float)) else 'STRING'
    item={'name':col,'role':role,'type':ftype}
    if col.endswith('_PCT') or col.endswith('_RH_PCT'): item['unit']='%'
    elif col.endswith('_MW_CM2'): item['unit']='mW/cm2'
    elif col.endswith('_MJ_CM2'): item['unit']='mJ/cm2'
    elif col.endswith('_UM'): item['unit']='um'
    elif col.endswith('_C'): item['unit']='degC'
    feature_schema.append(item)

manifest={
    'snapshot_schema_version':'1.0','snapshot_id':SNAPSHOT_ID,'task_profile_code':mod.TASK_PROFILE_CODE,'task_profile_version':mod.TASK_PROFILE_VERSION,
    'snapshot_purpose':'DEVELOPMENT','data_nature':'SYNTHETIC','production_eligible':False,'generated_at_utc':mod.FIXED_GENERATED_AT,
    'generator':{'name':'jsd-t07-uvpu-synthetic-shared-process-generator','version':GEN_VERSION,'seed':SEED,'row_count':400,'report_sheet_count':67},
    'source_basis':{
        'semantic_basis':'Original application test report semantics: one report sheet has one public application/cure context and multiple formulation columns.',
        'original_report_rule':'Rows 2 and 5-7 are public/shared conditions. When normalized, the same values are copied onto every experiment column in that sheet.',
        'source_excel':OUT_XLSX.name,
        'important_note':'All rows are synthetic and development-only.'
    },
    'shared_process_semantics':{
        'group_key':'source_sheet',
        'shared_fields':['X_PROCESS__COATING_SOLIDS_PCT','X_PROCESS__UVA_INTENSITY_MW_CM2','X_PROCESS__UV_ENERGY_MJ_CM2','X_ENV__TEMPERATURE_C','X_ENV__HUMIDITY_RH_PCT','X_CONTEXT__SUBSTRATE','X_CONTEXT__APPLICATION_METHOD','X_CONTEXT__CURING_SOURCE'],
        'per_experiment_fields':['formula composition','X_PROCESS__FILM_THICKNESS_UM','targets'],
        'rule':'All experiments on the same APP-xxx source sheet have identical shared public process/context X values.'
    },
    'formula_rules':{'basis':'PERCENT_SUM_100','sum_tolerance':0.02,'one_main_resin_per_experiment':True,'materials':[ *mod.RESINS,'DSP-3315','SS059/PC=1/1','S-48']},
    'feature_schema':feature_schema,'targets':mod.TARGETS,'target_statistics':target_stats,
    'missing_strategy':{'numeric_targets':'IEEE_NAN','note':'NaN means not tested; never convert to zero.'},
    'validation_strategy_hint':{'group_key':'formula_lineage_group','rule':'A formula lineage group must not be split across train and validation folds.'},
    'artifacts':{}
}
manifest['artifacts']={
    'measurements.parquet':{'sha256':mod.sha256_file(meas_path),'rows':400,'columns':len(mod.MEASUREMENT_COLUMNS),'format':'PARQUET'},
    'source-map.parquet':{'sha256':mod.sha256_file(src_path),'rows':400,'columns':len(mod.SOURCE_MAP_COLUMNS),'format':'PARQUET'}
}
(SNAP_DIR/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')
readme=f'''# UVPU_APPLICATION_FORMULATION Synthetic TrainingSnapshot V2 — Shared Process\n\n- 400 SYNTHETIC experiments\n- 67 source report sheets\n- fixed seed: {SEED}\n- development only; production_eligible=false\n\n## Key semantic rule\n\nThe original application report uses **public/shared process + multiple formulations**. For each APP-xxx sheet, rows 2 and 5–7 define shared temperature, humidity, substrate, coating method and UV cure conditions. During normalization, those shared conditions are copied onto every experiment in D:I, so every model row has its own complete X even though the raw Excel stores the process only once.\n\nDifferent APP-xxx sheets intentionally use different public process conditions. Formula composition and measured film thickness remain experiment-level fields.\n\n`/（未测试）` in the Excel and IEEE NaN in the snapshot both mean missing/not tested, never zero.\n'''
(SNAP_DIR/'README.md').write_text(readme,encoding='utf-8')
shutil.copy2(Path(__file__),SNAP_DIR/'create_shared_process_synthetic_v3.py')

# ---------- Verification ----------
wbv=load_workbook(OUT_XLSX,data_only=False)
assert len(wbv.sheetnames)==67
ids=[]
process_signatures=[]
for sidx,ws in enumerate(wbv.worksheets,1):
    expected=6 if sidx<67 else 4
    # B7 must be one exact condition, not a range.
    assert '–' not in str(ws['B7'].value) and '范围' not in str(ws['B7'].value)
    process_signatures.append((ws['G2'].value,ws['I2'].value,ws['B7'].value,ws['D27'].value))
    # Shared coating solids repeated across active experiment columns.
    solids=[]
    for c in range(4,4+expected):
        exp=ws.cell(16,c).value; assert exp and exp.startswith('SYN-UVPU-APP-'); ids.append(exp)
        solids.append(ws.cell(27,c).value)
        vals=[float(ws.cell(rr,c).value) for rr in range(17,26)]
        assert abs(sum(vals)-100.0)<0.02
        assert ws.cell(26,c).value.startswith('=IF(COUNT(')
        # no large required blanks
        for rr in range(8,16): assert ws.cell(rr,c).value not in (None,'')
        for rr in list(range(17,26))+list(range(27,37)): assert ws.cell(rr,c).value not in (None,'')
    assert len(set(solids))==1
assert len(ids)==400 and len(set(ids))==400
# Ensure process varies across reports.
assert len(set(process_signatures)) > 20

# Verify snapshot rows on each source sheet share public process fields.
shared=['X_PROCESS__COATING_SOLIDS_PCT','X_PROCESS__UVA_INTENSITY_MW_CM2','X_PROCESS__UV_ENERGY_MJ_CM2','X_ENV__TEMPERATURE_C','X_ENV__HUMIDITY_RH_PCT','X_CONTEXT__SUBSTRATE','X_CONTEXT__APPLICATION_METHOD','X_CONTEXT__CURING_SOURCE']
by_sheet={}
for m,s in records:
    by_sheet.setdefault(s['source_sheet'],[]).append(m)
for sheet,rows in by_sheet.items():
    for f in shared:
        assert len({r[f] for r in rows})==1,(sheet,f)

# zip excel + aligned snapshot
if OUT_ZIP.exists(): OUT_ZIP.unlink()
with zipfile.ZipFile(OUT_ZIP,'w',zipfile.ZIP_DEFLATED) as z:
    z.write(OUT_XLSX,OUT_XLSX.name)
    for p in sorted(SNAP_DIR.iterdir()):
        z.write(p,f'{SNAP_DIR.name}/{p.name}')

print('XLSX',OUT_XLSX,OUT_XLSX.stat().st_size)
print('SNAP',SNAP_DIR)
print('ZIP',OUT_ZIP,OUT_ZIP.stat().st_size)
print('reports',len(chunks),'experiments',len(ids),'unique_process_signatures',len(set(process_signatures)))
print('first 5 process rows:')
for ws in wbv.worksheets[:5]: print(ws.title,ws['G2'].value,ws['I2'].value,ws['B7'].value,ws['D27'].value)
