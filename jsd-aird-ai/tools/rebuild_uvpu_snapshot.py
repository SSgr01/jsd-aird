"""Rebuild the development UV/PU snapshot as standards-compliant Parquet.

The supplied V2 files were produced by a minimal custom writer. Their footer reports
400 rows, but PyArrow 21 decodes zero rows. This tool leaves those source artifacts
untouched and reconstructs a sibling snapshot from the authoritative synthetic XLSX.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


MEASUREMENT_COLUMNS = [
    "experiment_version_id",
    "X_FORMULA__SJ_230_PCT",
    "X_FORMULA__SJ_231_PCT",
    "X_FORMULA__SJ_232_PCT",
    "X_FORMULA__SJ_230_WASHED_PCT",
    "X_FORMULA__SJ_231_WASHED_PCT",
    "X_FORMULA__SJ_232_WASHED_PCT",
    "X_FORMULA__DSP_3315_PCT",
    "X_FORMULA__SS059_PC_1_1_PCT",
    "X_FORMULA__S_48_PCT",
    "X_PROCESS__COATING_SOLIDS_PCT",
    "X_PROCESS__UVA_INTENSITY_MW_CM2",
    "X_PROCESS__UV_ENERGY_MJ_CM2",
    "X_PROCESS__FILM_THICKNESS_UM",
    "X_ENV__TEMPERATURE_C",
    "X_ENV__HUMIDITY_RH_PCT",
    "X_CONTEXT__SUBSTRATE",
    "X_CONTEXT__APPLICATION_METHOD",
    "X_CONTEXT__CURING_SOURCE",
    "Y__WARPING_PET_INITIAL_CM",
    "Y__WARPING_PET_12H_CM",
    "Y__HARDNESS_PET_1KG_ORD",
    "Y__STEEL_WOOL_500G_CYCLES",
    "Y__STEEL_WOOL_1KG_CYCLES",
]

SOURCE_COLUMNS = [
    "experiment_version_id",
    "experiment_id",
    "project_id",
    "organization_id",
    "formula_lineage_group",
    "main_resin_code",
    "source_template_code",
    "source_file",
    "source_sheet",
    "source_range",
    "content_hash",
    "permission_scope",
    "data_nature",
    "snapshot_purpose",
    "generator_version",
    "generated_seed",
]

FORMULA_ROWS = {
    "X_FORMULA__SJ_230_PCT": (17, "SJ-230"),
    "X_FORMULA__SJ_231_PCT": (18, "SJ-231"),
    "X_FORMULA__SJ_232_PCT": (19, "SJ-232"),
    "X_FORMULA__SJ_230_WASHED_PCT": (20, "SJ-230-WASHED"),
    "X_FORMULA__SJ_231_WASHED_PCT": (21, "SJ-231-WASHED"),
    "X_FORMULA__SJ_232_WASHED_PCT": (22, "SJ-232-WASHED"),
    "X_FORMULA__DSP_3315_PCT": (23, None),
    "X_FORMULA__SS059_PC_1_1_PCT": (24, None),
    "X_FORMULA__S_48_PCT": (25, None),
}

MISSING_MARKERS = {"", "/", "/（未测试）", None}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while block := stream.read(1024 * 1024):
            digest.update(block)
    return digest.hexdigest()


def number(value: Any) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"-?\d+(?:\.\d+)?", str(value))
    if not match:
        raise ValueError(f"numeric value expected: {value!r}")
    return float(match.group())


def optional_number(value: Any, *, zero_text: str | None = None) -> float:
    if value in MISSING_MARKERS:
        return math.nan
    text = str(value)
    if zero_text and zero_text in text:
        return 0.0
    return number(value)


def hardness(value: Any) -> float:
    if value in MISSING_MARKERS:
        return math.nan
    match = re.match(r"\s*(?:(\d+)H|H)", str(value))
    if not match:
        raise ValueError(f"hardness value expected: {value!r}")
    return float(match.group(1) or 1)


def stable_json(value: Any) -> bytes:
    normalized = {
        key: (None if isinstance(item, float) and math.isnan(item) else item)
        for key, item in value.items()
    }
    return json.dumps(
        normalized,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        allow_nan=False,
    ).encode("utf-8")


def parse_sheet(sheet, source_name: str, seed: int) -> tuple[list[dict], list[dict]]:
    temperature = number(sheet.cell(2, 7).value)
    humidity = number(sheet.cell(2, 9).value)
    cure_text = str(sheet.cell(7, 2).value)
    intensity_match = re.search(r"(\d+(?:\.\d+)?)mW/cm", cure_text)
    energy_match = re.search(r"(\d+(?:\.\d+)?)mJ/cm", cure_text)
    if not intensity_match or not energy_match:
        raise ValueError(f"unable to parse UV conditions in {sheet.title}")
    intensity = float(intensity_match.group(1))
    energy = float(energy_match.group(1))

    measurements: list[dict] = []
    sources: list[dict] = []
    for column in range(4, sheet.max_column + 1):
        experiment_id = sheet.cell(16, column).value
        if not experiment_id:
            continue
        experiment_id = str(experiment_id).strip()
        version_id = f"{experiment_id}-V1"
        formula: dict[str, float] = {}
        active_resins: list[str] = []
        for feature, (row, resin_code) in FORMULA_ROWS.items():
            value = float(sheet.cell(row, column).value or 0.0)
            formula[feature] = value
            if resin_code and value > 1e-9:
                active_resins.append(resin_code)
        if len(active_resins) != 1:
            raise ValueError(f"{version_id} does not have exactly one main resin")

        solids = number(sheet.cell(27, column).value)
        if solids <= 1.0:
            solids *= 100.0
        measurement = {
            "experiment_version_id": version_id,
            **formula,
            "X_PROCESS__COATING_SOLIDS_PCT": solids,
            "X_PROCESS__UVA_INTENSITY_MW_CM2": intensity,
            "X_PROCESS__UV_ENERGY_MJ_CM2": energy,
            "X_PROCESS__FILM_THICKNESS_UM": number(sheet.cell(31, column).value),
            "X_ENV__TEMPERATURE_C": temperature,
            "X_ENV__HUMIDITY_RH_PCT": humidity,
            "X_CONTEXT__SUBSTRATE": "PET_100UM_OPTICAL",
            "X_CONTEXT__APPLICATION_METHOD": "WIRE_BAR_ROLL_COAT_40UM",
            "X_CONTEXT__CURING_SOURCE": "MERCURY_LAMP_UV",
            "Y__WARPING_PET_INITIAL_CM": optional_number(sheet.cell(32, column).value, zero_text="平整"),
            "Y__WARPING_PET_12H_CM": optional_number(sheet.cell(33, column).value, zero_text="平整"),
            "Y__HARDNESS_PET_1KG_ORD": hardness(sheet.cell(34, column).value),
            "Y__STEEL_WOOL_500G_CYCLES": optional_number(sheet.cell(35, column).value),
            "Y__STEEL_WOOL_1KG_CYCLES": optional_number(sheet.cell(36, column).value),
        }
        experiment_number = int(re.search(r"(\d+)$", experiment_id).group(1))
        source = {
            "experiment_version_id": version_id,
            "experiment_id": experiment_id,
            "project_id": "SYNTH-UVPU-APPLICATION-PROJECT",
            "organization_id": "SYNTHETIC-DEVELOPMENT",
            "formula_lineage_group": f"FLG-{((experiment_number - 1) // 10) + 1:03d}",
            "main_resin_code": active_resins[0],
            "source_template_code": "APPLICATION_TEST_REPORT_V1",
            "source_file": source_name,
            "source_sheet": sheet.title,
            "source_range": f"{sheet.cell(8, column).coordinate}:{sheet.cell(36, column).coordinate}",
            "content_hash": hashlib.sha256(stable_json(measurement)).hexdigest(),
            "permission_scope": "INTERNAL_TEST_ONLY",
            "data_nature": "SYNTHETIC",
            "snapshot_purpose": "DEVELOPMENT",
            "generator_version": "2.0.1-pyarrow-repair",
            "generated_seed": seed,
        }
        measurements.append(measurement)
        sources.append(source)
    return measurements, sources


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", required=True, type=Path)
    parser.add_argument("--source-manifest", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    args.output.mkdir(parents=True, exist_ok=True)
    source_manifest = json.loads(args.source_manifest.read_text(encoding="utf-8"))
    seed = int(source_manifest["generator"]["seed"])
    workbook = load_workbook(args.xlsx, read_only=True, data_only=True)
    measurements: list[dict] = []
    sources: list[dict] = []
    for sheet in workbook.worksheets:
        sheet_measurements, sheet_sources = parse_sheet(sheet, args.xlsx.name, seed)
        measurements.extend(sheet_measurements)
        sources.extend(sheet_sources)

    if len(measurements) != 400 or len(sources) != 400:
        raise ValueError("expected exactly 400 reconstructed experiments")

    measurements_frame = pd.DataFrame(measurements, columns=MEASUREMENT_COLUMNS)
    source_frame = pd.DataFrame(sources, columns=SOURCE_COLUMNS)
    measurements_path = args.output / "measurements.parquet"
    source_path = args.output / "source-map.parquet"
    measurements_frame.to_parquet(measurements_path, engine="pyarrow", compression="zstd", index=False)
    source_frame.to_parquet(source_path, engine="pyarrow", compression="zstd", index=False)

    manifest = dict(source_manifest)
    manifest["snapshot_id"] = f"{source_manifest['snapshot_id']}-PYARROW1"
    manifest["generated_at_utc"] = datetime.now(UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    manifest["generator"] = {
        **source_manifest["generator"],
        "version": "2.0.1-pyarrow-repair",
        "source_snapshot_id": source_manifest["snapshot_id"],
        "repair_reason": "Original custom Parquet files decode as zero rows in PyArrow 21",
    }
    target_statistics = {}
    for target_code in source_manifest["targets"]:
        values = measurements_frame[target_code]
        valid = values.dropna()
        target_statistics[target_code] = {
            "valid_count": int(valid.size),
            "missing_count": int(values.isna().sum()),
            "missing_rate": float(values.isna().mean()),
            "min": float(valid.min()),
            "max": float(valid.max()),
        }
    manifest["target_statistics"] = target_statistics
    manifest["artifacts"] = {
        "measurements.parquet": {
            "sha256": sha256_file(measurements_path),
            "rows": len(measurements_frame),
            "columns": len(measurements_frame.columns),
            "format": "PARQUET",
            "writer": "pyarrow",
        },
        "source-map.parquet": {
            "sha256": sha256_file(source_path),
            "rows": len(source_frame),
            "columns": len(source_frame.columns),
            "format": "PARQUET",
            "writer": "pyarrow",
        },
    }
    (args.output / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )
    measurements_frame.head(50).to_csv(args.output / "measurements_preview.csv", index=False)
    source_frame.head(50).to_csv(args.output / "source-map_preview.csv", index=False)
    (args.output / "README.md").write_text(
        "# PyArrow-compatible UV/PU synthetic snapshot\n\n"
        "This development-only snapshot was reconstructed from the supplied synthetic XLSX. "
        "The original custom Parquet artifacts are preserved in the sibling directory. "
        "Displayed XLSX values are rounded, so this dataset is suitable for T07-A workflow "
        "verification, not production-effect claims.\n",
        encoding="utf-8",
    )
    print(json.dumps({"rows": len(measurements_frame), "output": str(args.output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
