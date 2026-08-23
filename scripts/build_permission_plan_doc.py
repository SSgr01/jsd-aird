from pathlib import Path
from datetime import datetime

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"G:\Projects\jsd-aird")
OUT = ROOT / "docs" / "杰事达材料研发系统_系统权限管理功能开发步骤计划_V1.0.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
GREEN = "1F6B4F"
GOLD = "7A5A00"
RED = "9B1C1C"


def set_run_font(run, name="Microsoft YaHei", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_border(cell, color="D0D5DD", size="4"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("222222")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Microsoft YaHei"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)
    subtitle.paragraph_format.space_after = Pt(18)

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 11.5, DARK_BLUE, 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ["List Bullet", "List Number"]:
        style = styles[list_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    for name, size, color, bold in [
        ("Kicker", 10, BLUE, True),
        ("Small Text", 9, MUTED, False),
        ("Table Text", 9.2, "222222", False),
        ("Table Header", 9.2, NAVY, True),
        ("Callout Text", 10, NAVY, False),
    ]:
        if name in styles:
            style = styles[name]
        else:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = bold
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def configure_header_footer(section):
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("杰事达材料研发系统  |  系统权限管理功能开发步骤计划")
    set_run_font(r, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    r = fp.add_run("内部项目计划  |  第 ")
    set_run_font(r, size=8.5, color=MUTED)
    add_page_field(fp)
    r = fp.add_run(" 页")
    set_run_font(r, size=8.5, color=MUTED)


def add_paragraph(doc, text="", style=None, bold_prefix=None, color=None, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True, color=color or "222222")
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, color=color or "222222")
    else:
        r = p.add_run(text)
        set_run_font(r, color=color or "222222")
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_run_font(r, size=10.5, color="222222")
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(text)
    set_run_font(r, size=10.5, color="222222")
    return p


def add_callout(doc, label, text, fill=CALLOUT, label_color=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360], indent_dxa=120)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color="D0D5DD", size="6")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + "  ")
    set_run_font(r, size=9.5, color=label_color, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=9.5, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths, font_size=9.0, header_fill=LIGHT_BLUE, indent=120):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_geometry(table, widths, indent_dxa=indent)
    header = table.rows[0]
    repeat_table_header(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, header_fill)
        set_cell_border(cell)
        p = cell.paragraphs[0]
        p.style = "Table Header"
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(str(value))
        set_run_font(r, size=font_size, color=NAVY, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_border(cell)
            if isinstance(value, tuple):
                text, fill = value
                set_cell_shading(cell, fill)
            else:
                text = value
            parts = str(text).split("\n")
            for j, part in enumerate(parts):
                p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
                p.style = "Table Text"
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                r = p.add_run(part)
                set_run_font(r, size=font_size, color="222222")
    set_table_geometry(table, widths, indent_dxa=indent)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_stage(doc, stage_no, title, objective, duration, owner, dependencies, tasks, outputs, gate, effort):
    doc.add_heading(f"阶段 {stage_no}: {title}", level=2)
    add_paragraph(doc, f"阶段目标：{objective}")
    add_table(
        doc,
        ["计划周期", "主要负责人", "前置依赖", "估算工作量"],
        [[duration, owner, dependencies, effort]],
        [1350, 2500, 2900, 2610],
        font_size=8.8,
        header_fill=LIGHT_GRAY,
    )
    doc.add_heading("功能开发步骤", level=3)
    for task in tasks:
        add_number(doc, task)
    doc.add_heading("阶段交付物", level=3)
    for item in outputs:
        add_bullet(doc, item)
    add_callout(doc, "阶段门禁", gate, fill="EFF6FF", label_color=BLUE)


def add_source_note(doc, text):
    p = doc.add_paragraph(style="Small Text")
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("依据：" + text)
    set_run_font(r, size=8.5, color=MUTED, italic=True)


def build_doc():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    configure_section(section)
    configure_header_footer(section)

    # Cover / customer-pack style opening block.
    p = doc.add_paragraph(style="Kicker")
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run("功能开发步骤计划 | V1.0")
    set_run_font(r, size=10, color=BLUE, bold=True)

    p = doc.add_paragraph(style="Title")
    r = p.add_run("杰事达材料研发系统")
    set_run_font(r, size=25, color=NAVY, bold=True)
    p = doc.add_paragraph(style="Subtitle")
    r = p.add_run("系统权限管理功能开发步骤计划")
    set_run_font(r, size=15, color=MUTED)

    add_paragraph(doc, "将统一权限从部门级粗粒度控制落地为：角色默认基线 + 个人权限覆盖 + 数据范围过滤 + 审计可追溯。", color=NAVY)

    add_table(
        doc,
        ["项目属性", "计划内容", "项目属性", "计划内容"],
        [
            ["适用范围", "一期统一权限管理功能", "计划版本", "V1.0"],
            ["建议周期", "8 周日历周期", "估算工作量", "约 75 人日"],
            ["交付对象", "产品、研发、测试、实施、运维", "发布策略", "灰度上线 + 可回滚"],
        ],
        [1350, 3330, 1350, 3330],
        font_size=9.2,
        header_fill=LIGHT_BLUE,
    )

    add_callout(doc, "推荐排期", "采用 2 名后端、1 名前端、1 名测试，产品/架构与运维共享投入的配置，按 8 个阶段推进。后端可并行开展数据迁移、认证和权限核心能力，但阶段门禁必须按顺序通过。", fill="EAF4EF", label_color=GREEN)

    add_paragraph(doc, "基线文档", style="Heading 2")
    add_bullet(doc, "《杰事达材料研发系统_系统权限管理功能需求规格说明书_V1.0》")
    add_bullet(doc, "《杰事达材料研发系统_系统权限管理详细设计说明书_V1.0》")

    doc.add_page_break()

    doc.add_heading("1. 计划总览", level=1)
    add_paragraph(doc, "本计划将需求规格书中的用户、角色、权限目录、个人数据范围、AI 与敏感数据安全、操作日志和非功能要求，映射为可独立开发、联调和验收的阶段。详细设计书中的数据库迁移、AuthorizationService、DataScopeFilter、REST 接口、缓存、审计和灰度回滚要求，作为各阶段的实现约束。")
    add_callout(doc, "执行原则", "先冻结权限契约，再建设判定内核；先保证后端强制执行，再补齐前端菜单和按钮控制；任何列表/搜索必须在 Repository 查询层应用 DataScopeFilter，不能查全量后在前端过滤。", fill="FFF8E8", label_color=GOLD)

    doc.add_heading("1.1 一期交付边界", level=2)
    for item in [
        "账号与会话：本地账号登录、退出、当前用户、修改密码、会话撤销、禁用账号即时失效、登录失败限流和安全响应头。",
        "授权模型：固定权限目录、内置角色、单主角色、个人权限覆盖、有效权限计算、默认拒绝和权限变更版本。",
        "数据范围：全部、本人、本人及下级、指定部门、指定项目/范围，以及业务模块按资源归属过滤。",
        "管理端：用户管理、角色管理、权限配置、个人权限差异/覆盖、数据范围配置、审计查询。",
        "安全与运维：AI 外部处理前的敏感数据权限门禁、操作审计、缓存失效、指标、迁移、灰度和回滚。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("1.2 一期非目标", level=2)
    for item in [
        "不在本期引入完整 SSO、SCIM、LDAP/AD 同步等外部身份目录能力；通过接口抽象为后续扩展点。",
        "不允许管理员在页面自由创建权限编码；权限定义由各模块通过 PermissionDefinitionRegistry 注册并幂等同步。",
        "不以部门字段直接替代数据范围；组织关系只作为范围解析输入，最终必须落到目标资源归属判断。",
        "不把前端隐藏菜单、路由守卫或请求头作为唯一安全边界；后端 AuthorizationService 是强制执行点。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("1.3 计划假设", level=2)
    add_table(
        doc,
        ["假设项", "本计划采用的判断"],
        [
            ["团队配置", "2 后端 + 1 前端 + 1 测试；产品/架构和运维各按 0.5 人投入。"],
            ["基础设施", "现有 Java 服务、数据库、缓存、日志和部署流水线可复用；不把基础设施重构计入本期。"],
            ["组织与资源", "现有用户/组织/部门主数据存在；各业务模块需补齐资源归属字段或查询适配。"],
            ["权限矩阵", "需求附录 A 的内置角色及默认矩阵为种子数据来源，最终以评审冻结版为准。"],
            ["数据风险", "AI 外部处理场景由业务模块显式标记资源是否含敏感数据，并提交用途、提供方、资源类型和摘要。"],
        ],
        [2200, 7160],
        font_size=9.2,
    )

    doc.add_heading("2. 阶段化功能开发步骤", level=1)
    add_paragraph(doc, "每个阶段均设置可核验的交付物和阶段门禁。未通过门禁的阶段，不进入依赖它的下一个功能批次；允许同一阶段内部并行，但不允许绕过后端授权、数据范围和安全测试。")

    add_stage(
        doc,
        "0",
        "需求基线与权限契约冻结",
        "把文档中的业务规则转化为研发、测试和实施可以共同使用的编码、接口、数据范围和验收基线。",
        "第 1 周前 2-3 天",
        "产品负责人 / 技术负责人",
        "无",
        [
            "建立权限目录清单：模块编码、动作编码、权限编码、资源类型、是否涉及敏感数据、对应后端执行点。",
            "确认内置角色、角色默认权限、个人覆盖权限的优先级，以及恢复角色默认权限的行为。",
            "冻结数据范围枚举和目标资源归属规则，明确各模块的部门、项目、负责人等关联字段。",
            "冻结认证接口、IAM 管理接口、错误码、审计字段和前端需要的权限快照结构。",
            "把验收场景整理为可执行测试账号矩阵：超级管理员、系统管理员、普通用户、跨部门用户、禁用用户和无权限用户。",
        ],
        [
            "权限目录/编码清单 V1.0",
            "角色默认矩阵和覆盖规则确认单",
            "数据范围映射表及业务模块接入清单",
            "接口契约、错误码和验收账号矩阵",
        ],
        "所有权限动作均有固定编码、后端执行点和验收用例；每个需要数据范围的列表/详情/导出接口都有资源归属字段或适配方案。",
        "3-5 人日",
    )

    add_stage(
        doc,
        "1",
        "数据模型与数据库迁移",
        "落地用户、角色、权限、覆盖、数据范围、会话和审计相关表结构，并保证新老数据可升级、可验证、可回滚。",
        "第 1 周",
        "后端 1 / DBA / 运维",
        "阶段 0 的模型冻结",
        [
            "创建或扩展用户、角色、权限定义、角色权限绑定、个人权限覆盖、数据范围目标、会话、登录失败记录等表。",
            "为用户状态、角色关系、权限编码、主体-目标关系和权限版本建立唯一约束、索引及必要的审计字段。",
            "实现 PermissionDefinitionRegistry 的幂等同步：服务启动时注册，IAM 侧可查询，版本升级不产生重复权限。",
            "准备内置角色和默认权限种子数据；明确已有用户的角色映射、默认范围和异常数据处理方式。",
            "在测试库完成全新安装、从现网结构升级、重复执行迁移和回滚演练。",
        ],
        [
            "数据库迁移脚本、索引和约束",
            "权限目录初始化/幂等同步实现",
            "内置角色及默认矩阵种子数据",
            "升级、重复执行和回滚记录",
        ],
        "新库可初始化，老库可升级；迁移可重复执行且不丢失原有组织、用户、审计和业务数据；种子数据与冻结矩阵一致。",
        "6-8 人日",
    )

    add_stage(
        doc,
        "2",
        "登录、会话与账户安全",
        "建立可信 ActorContext 和会话生命周期，为后续所有授权判定提供稳定的身份输入。",
        "第 2 周",
        "后端 1 / 前端",
        "阶段 1 的用户和会话表",
        [
            "实现登录、退出、当前用户、修改密码、会话列表、撤销会话、管理员重置密码等接口。",
            "使用现有安全密码哈希策略；登录失败按账号/IP/时间窗口限流，并避免返回可枚举用户是否存在的错误信息。",
            "登录成功建立 HttpOnly、Secure、SameSite 会话；服务端只信任会话中的用户身份，不接受前端自报用户 ID、角色或部门。",
            "账号禁用、密码重置、角色/权限重大变更时递增 authVersion 或撤销会话，使旧会话不能继续使用。",
            "补齐生产环境安全响应头、错误码和前端登录态处理；为未来 SSO/LDAP 保留认证适配接口。",
        ],
        [
            "认证 REST 接口和会话服务",
            "登录页、退出和登录态恢复",
            "限流、密码策略、安全响应头",
            "会话撤销及禁用账号失效测试",
        ],
        "禁用用户的现有会话立即失效；伪造请求头不能切换用户或提升权限；登录失败限流和敏感操作审计可被测试验证。",
        "7-9 人日",
    )

    add_stage(
        doc,
        "3",
        "权限目录、授权判定与数据范围内核",
        "实现一期最核心的 AuthorizationService、有效权限算法和 DataScopeFilter，并成为业务模块唯一的后端授权入口。",
        "第 2-3 周",
        "后端 2 / 技术负责人",
        "阶段 1 的身份与权限数据模型",
        [
            "定义 Subject、PermissionCode、ResourceType、Action、ScopeType、DataScopeFilter 和 AuthorizationDecision 等公共类型。",
            "实现判定顺序：账号状态 -> 权限目录 -> 个人覆盖；个人显式覆盖优先于角色默认；未配置时按角色默认；最终拒绝优先。",
            "实现全部、本人、本人及下级、指定部门、指定项目/范围等范围解析，并在目标资源不属于当前组织时返回 DATA_SCOPE_DENIED。",
            "实现“允许访问”与“可见数据范围”分离：动作授权通过后，列表/搜索/导出仍必须应用 DataScopeFilter。",
            "实现默认拒绝、权限版本、缓存失效和高风险允许/拒绝审计；禁止 Controller 直接读权限表或直接查库判权。",
            "为每类范围、覆盖恢复、角色变更、部门变更、跨组织资源和缓存过期编写单元测试。",
        ],
        [
            "IAM 公共接口和授权判定实现",
            "DataScopeResolver / DataScopeFilter",
            "错误码、缓存键/版本和审计事件钩子",
            "授权算法单元测试和边界测试",
        ],
        "角色默认、个人允许、个人拒绝、恢复默认和未配置场景均符合规则；业务 Repository 能收到范围过滤条件；默认拒绝和跨组织拒绝测试通过。",
        "10-14 人日",
    )

    add_stage(
        doc,
        "4",
        "IAM 管理接口与前端管理页面",
        "让管理员可以管理账号、角色、权限覆盖和数据范围，并让前端路由/菜单/按钮与后端权限快照保持一致。",
        "第 3-4 周",
        "后端 1 / 前端 1",
        "阶段 3 的公共授权服务和接口契约",
        [
            "实现用户查询、创建/编辑、启用/禁用、密码重置、会话查看和会话撤销接口。",
            "实现角色查询、角色权限查看/编辑、权限目录查询、个人权限覆盖、恢复角色默认权限和有效权限差异接口。",
            "实现数据范围目标的查询、保存和删除；保存时校验目标组织归属、资源类型和重复关系。",
            "实现用户管理、角色管理、权限配置、个人差异和数据范围页面；权限目录使用固定编码，只读展示模块注册结果。",
            "实现前端登录态、权限快照、路由守卫、菜单过滤和按钮级控制；所有管理操作仍由后端二次判权。",
            "统一无权限、超范围、账号禁用、会话过期和参数错误的提示，不泄露资源是否存在等敏感信息。",
        ],
        [
            "IAM 管理 REST 接口",
            "账号/角色/权限/数据范围管理页面",
            "有效权限差异和恢复默认交互",
            "前端权限快照、路由和按钮控制",
        ],
        "管理员可以完成一条完整配置链路：创建用户 -> 分配角色 -> 设置个人覆盖 -> 设置范围 -> 查询有效权限；越权请求在浏览器直接调用 API 时仍被拒绝。",
        "10-12 人日",
    )

    add_stage(
        doc,
        "5",
        "业务模块接入与页面控制",
        "把权限内核嵌入实际业务读写链路，确保数据范围在查询层生效，而不是停留在 IAM 管理页面。",
        "第 5-6 周",
        "后端 2 / 前端 1 / 各模块负责人",
        "阶段 3 完成；阶段 4 提供稳定接口",
        [
            "按优先级分波次接入：先接入核心研发/材料数据列表与详情，再接入创建、编辑、删除、导出和批量操作。",
            "每个模块在应用服务调用 AuthorizationService，Controller 不直接解析权限表；列表、搜索、导出统一把 DataScopeFilter 传入 Repository。",
            "为目标资源补齐归属信息：部门、负责人、项目、创建人或指定范围；缺少归属时先阻断上线并补适配。",
            "接入页面路由、菜单、按钮、批量操作和导出权限；前端隐藏只作为体验优化，不能替代后端判定。",
            "按模块建立正向、反向、跨部门、跨项目、无范围和默认拒绝用例，并记录 SQL/查询条件确实带有范围过滤。",
            "形成业务模块接入清单和代码审查规则，禁止出现“先查全量再在前端过滤”以及绕过公共授权服务的实现。",
        ],
        [
            "核心业务模块接入 PR/代码审查记录",
            "列表/详情/写操作/导出授权和范围过滤",
            "模块级权限与数据范围回归用例",
            "资源归属字段和查询适配说明",
        ],
        "每个纳入一期的业务模块都能从登录用户到数据库查询完成闭环；跨部门/跨项目数据不会通过列表、详情、导出或批量接口泄露。",
        "12-16 人日",
    )

    add_stage(
        doc,
        "6",
        "AI 敏感数据、审计、缓存与可观测性",
        "把高风险 AI 外发、权限变化和管理员操作纳入统一安全控制和可追溯链路。",
        "第 6-7 周",
        "后端 1 / 安全 / 运维",
        "阶段 3 的判定、审计钩子和缓存版本",
        [
            "在 AI 外部处理入口前执行权限判定；当资源可能含敏感数据时，校验数据类型、用途、提供方、资源类型和摘要。",
            "区分 AI 使用权限和业务资源数据范围权限：具备 AI 动作权限不代表可以读取超出个人数据范围的资源。",
            "记录登录、登出、失败、用户/角色/权限/范围变更、敏感 AI 操作、授权拒绝、高风险允许和会话撤销审计。",
            "实现权限缓存读取、按主体/权限/版本失效、角色变更和个人覆盖变更后的即时刷新；为缓存命中率和拒绝量提供指标。",
            "补齐结构化日志、traceId、subjectId、permissionCode、resourceType、resourceId、decision、reason 和 clientIp 等字段。",
            "对审计内容做最小化和脱敏处理，避免把原始敏感材料写入日志或发送至不受控的外部系统。",
        ],
        [
            "AI 外发前置门禁和拒绝提示",
            "审计事件模型、查询接口和日志字段",
            "缓存失效、指标和告警规则",
            "敏感数据不外泄验证记录",
        ],
        "高风险 AI 操作无权限或超数据范围时被阻断；所有权限变更与敏感操作可按主体、动作、资源和时间检索；缓存不会造成权限旧值长期生效。",
        "7-9 人日",
    )

    add_stage(
        doc,
        "7",
        "联调、验收、灰度上线与回滚",
        "把功能、性能、安全、迁移和运维证据汇总为可发布版本，并通过灰度和回滚演练降低上线风险。",
        "第 7-8 周",
        "测试负责人 / 项目经理 / 运维",
        "阶段 4-6 完成，业务模块接入清单冻结",
        [
            "执行后端单元测试、接口集成测试、前端组件/页面测试和跨模块端到端测试。",
            "执行权限绕过专项测试：伪造请求头、篡改用户/角色/部门、直接调用隐藏 API、分页/导出/批量接口和资源 ID 越权。",
            "执行数据范围专项测试：全部、本人、本人及下级、指定部门、指定项目、无范围、跨组织、资源归属缺失。",
            "执行迁移、初始化、权限缓存刷新、会话撤销、禁用账号、错误码、审计检索和告警验证。",
            "在预生产按灰度方案发布，先观察登录失败、拒绝率、范围过滤异常、接口时延、缓存命中和审计写入。",
            "演练回滚：关闭新权限判定开关/恢复上一版本、保留审计、恢复数据迁移前快照，并验证业务可用性。",
        ],
        [
            "测试报告、缺陷清单和回归结果",
            "安全测试报告和越权测试证据",
            "迁移/灰度/回滚演练记录",
            "上线检查单、运维监控和应急联系人",
        ],
        "需求验收场景全部通过；P0/P1 缺陷清零；关键接口无越权和范围泄露；灰度指标稳定；回滚在预生产可重复执行。",
        "10-12 人日",
    )

    doc.add_heading("3. 工作包与人员分工", level=1)
    add_paragraph(doc, "工作包按可分配、可审查和可验收拆分。建议用工作包编号建立 Jira/禅道任务，并在合并请求、测试用例和发布单中回填编号。")
    add_table(
        doc,
        ["工作包", "负责人", "主要内容", "完成判定"],
        [
            ["WP-01", "产品/架构", "权限编码、角色矩阵、范围规则、错误码和验收账号", "评审冻结，需求与设计无未决冲突"],
            ["WP-02", "后端/DBA", "IAM 表结构、索引、迁移、种子数据", "新库/老库/重复迁移/回滚均通过"],
            ["WP-03", "后端 1", "登录、密码、会话、限流、禁用失效", "认证和会话安全用例通过"],
            ["WP-04", "后端 2", "权限目录注册、同步和版本管理", "模块权限可幂等注册和查询"],
            ["WP-05", "后端 2", "AuthorizationService 与判定算法", "覆盖/默认拒绝/审计决策测试通过"],
            ["WP-06", "后端 2", "DataScopeResolver、Filter 和 Repository 适配", "所有列表/搜索/导出均带范围条件"],
            ["WP-07", "后端 1", "用户/角色/权限/范围管理 API", "管理员完整配置链路可用"],
            ["WP-08", "前端", "登录态、管理页、路由/菜单/按钮控制", "前端体验与后端判定一致"],
            ["WP-09", "各模块负责人", "核心业务模块读写、导出和批量操作接入", "模块接入清单与代码审查通过"],
            ["WP-10", "安全/后端", "AI 外发门禁、敏感字段和错误提示", "无权限/超范围均被阻断"],
            ["WP-11", "后端/运维", "审计、缓存、指标、日志和告警", "事件可查，缓存可失效，指标可观测"],
            ["WP-12", "测试/运维", "集成、安全、性能、迁移、灰度和回滚", "发布门禁全部通过"],
        ],
        [1050, 1450, 4550, 2310],
        font_size=8.7,
    )

    doc.add_heading("4. 依赖关系与推荐排期", level=1)
    add_paragraph(doc, "推荐采用“核心内核先行、管理端并行、业务模块分波次接入”的排期。阶段 2 可与阶段 1 后半段并行，阶段 4 可与阶段 3 后半段并行，但阶段 5 必须以阶段 3 的授权内核稳定为前提。")
    add_table(
        doc,
        ["周次", "主线目标", "可并行工作", "周末应达到的状态"],
        [
            ["W1", "阶段 0 + 阶段 1", "迁移脚本、权限目录梳理", "权限契约冻结，迁移可在测试库执行"],
            ["W2", "阶段 2 + 阶段 3 启动", "登录页面、判定公共类型", "可信会话可用，授权算法骨架完成"],
            ["W3", "阶段 3", "IAM API 设计、核心单元测试", "授权、范围、错误码和缓存版本可测"],
            ["W4", "阶段 4", "第一批业务模块接入准备", "管理员可配置有效权限和数据范围"],
            ["W5", "阶段 5 第一波", "前端菜单/按钮、模块查询适配", "核心列表/详情完成后端判权与范围过滤"],
            ["W6", "阶段 5 第二波 + 阶段 6", "AI 门禁、审计、缓存指标", "主要业务写操作/导出纳入控制"],
            ["W7", "阶段 6 收尾 + 阶段 7", "安全专项、迁移和灰度演练", "测试报告和上线候选版本形成"],
            ["W8", "阶段 7", "用户验收、发布和回滚演练", "灰度稳定，发布门禁通过，具备回滚路径"],
        ],
        [900, 2450, 3000, 3010],
        font_size=8.8,
    )

    doc.add_heading("5. 测试与验收门禁", level=1)
    add_paragraph(doc, "测试不只验证页面是否隐藏按钮，还要验证后端授权决策、SQL/Repository 范围过滤、会话状态、缓存版本和审计证据。")
    add_table(
        doc,
        ["验收主题", "必须验证的场景", "证据"],
        [
            ["有效权限计算", "无个人覆盖时继承角色默认；个人允许/拒绝优先；恢复后回到角色默认。", "授权单元测试 + API 测试"],
            ["账号与会话", "禁用账号现有会话失效；密码重置/权限重大变更可撤销会话；超时返回统一错误。", "安全测试报告 + 会话日志"],
            ["后端强制执行", "伪造用户/角色/部门请求头、直接调用隐藏 API、篡改资源 ID 均不能越权。", "越权测试记录"],
            ["数据范围", "全部、本人、本人及下级、部门、项目、无范围和跨组织资源均符合规则。", "查询结果 + SQL/Repository 断言"],
            ["AI 敏感数据", "缺少 AI 动作权限、超资源范围、敏感字段外发条件不完整时被阻断。", "门禁日志 + 脱敏检查"],
            ["审计与观测", "登录失败、权限变更、拒绝、敏感 AI 操作和会话撤销可按条件检索。", "审计查询截图/接口响应"],
            ["迁移与回滚", "新旧库升级、重复执行、种子数据、灰度、关闭开关和回滚后业务可用。", "演练记录 + 发布检查单"],
        ],
        [1700, 5050, 2610],
        font_size=8.7,
    )

    doc.add_heading("6. 风险、决策点与应对", level=1)
    add_table(
        doc,
        ["风险/决策点", "影响", "应对措施", "责任人"],
        [
            ["业务资源缺少明确归属", "无法可靠执行部门/项目/本人范围，可能导致过度开放或误拒绝。", "阶段 0 冻结归属字段；阶段 5 未补齐的接口不得上线。", "模块负责人"],
            ["权限编码与后端执行点脱节", "出现前端有开关、后端无判定的幽灵权限。", "权限只由 Registry 注册；每个编码必须绑定执行点和测试用例。", "架构/后端"],
            ["个人覆盖复杂度失控", "管理员难以解释用户为什么拥有/失去权限。", "保留单主角色；提供有效权限差异、来源和恢复默认能力。", "产品/IAM"],
            ["前端过滤代替后端过滤", "列表、导出或批量接口发生数据泄露。", "代码审查禁止全量查询；Repository 必须接收 DataScopeFilter。", "后端负责人"],
            ["缓存旧值导致权限延迟生效", "禁用、撤权或范围变更后仍能访问。", "主体/权限/版本缓存键 + 变更失效 + 关键操作回源校验。", "后端/运维"],
            ["AI 外发边界不清", "敏感材料被错误发送给外部提供方。", "业务显式标记敏感性；外发前检查动作权限、资源范围、用途和提供方。", "安全/产品"],
            ["存量用户映射不完整", "上线后出现大量无法登录、权限过宽或业务中断。", "先做数据盘点和影子判定；灰度观察；保留可回滚迁移。", "项目经理/DBA"],
        ],
        [1900, 2450, 3650, 1360],
        font_size=8.5,
    )

    doc.add_heading("7. Definition of Done 与发布检查单", level=1)
    doc.add_heading("7.1 功能完成标准", level=2)
    for item in [
        "需求规格书第 5-11 章范围内的功能均有对应工作包、代码、测试用例和验收证据。",
        "详细设计书第 3-10 章的实体、接口、判定算法、数据范围、缓存、审计和前端状态均已落地或在发布单中明确未纳入项。",
        "所有新增权限编码有模块注册、后端执行点、默认矩阵和前端展示/控制映射。",
        "关键列表、详情、导出、批量和 AI 操作均经过后端授权与数据范围过滤。",
        "无 P0/P1 缺陷；P2 缺陷有产品确认、临时措施和后续计划。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("7.2 上线前检查顺序", level=2)
    for item in [
        "冻结权限目录、角色矩阵、范围映射、迁移版本和发布版本。",
        "备份数据库和权限配置；在预生产完成迁移、重复迁移和回滚演练。",
        "验证超级管理员应急账号、普通用户测试账号、禁用账号和无权限账号。",
        "执行安全专项：请求头伪造、资源 ID 越权、分页/导出/批量、缓存旧值和 AI 外发。",
        "确认审计、告警、接口时延、拒绝率、缓存命中率和错误码监控已接入。",
        "灰度发布并观察至少一个完整业务高峰窗口；异常时按回滚剧本执行并保留审计。",
    ]:
        add_number(doc, item)

    doc.add_heading("8. 交付物清单", level=1)
    add_table(
        doc,
        ["交付物", "产生阶段", "用途"],
        [
            ["权限目录/编码清单", "阶段 0", "产品、开发、测试共用的固定权限契约"],
            ["角色默认矩阵与覆盖规则", "阶段 0/1", "种子数据、管理页面和验收基线"],
            ["数据范围映射表", "阶段 0/5", "业务资源归属和 Repository 适配依据"],
            ["数据库迁移与回滚脚本", "阶段 1/7", "版本升级、灰度和应急恢复"],
            ["IAM API 与公共接口说明", "阶段 2-4", "前后端和业务模块接入契约"],
            ["模块接入清单与代码审查记录", "阶段 5", "证明业务链路未绕过授权和范围过滤"],
            ["测试报告与安全测试证据", "阶段 7", "验收、上线和问题追踪"],
            ["灰度发布、监控和回滚手册", "阶段 6-7", "上线运行和故障处置"],
        ],
        [2600, 1800, 4960],
        font_size=9.0,
    )

    doc.add_heading("附录 A. 需求与设计章节映射", level=1)
    add_table(
        doc,
        ["来源文档章节", "计划阶段", "落地内容"],
        [
            ["需求第 3-4 章：范围、原则、术语与有效权限规则", "阶段 0/3", "权限编码、判定优先级、默认拒绝和范围模型"],
            ["需求第 5-6 章：用户、角色、权限目录和动作", "阶段 1/3/4", "实体、种子数据、Registry、管理 API 和页面"],
            ["需求第 7 章：个人数据范围", "阶段 0/3/5", "范围解析、过滤器、资源归属和模块接入"],
            ["需求第 8-9 章：AI 安全、操作日志与审计", "阶段 6", "AI 外发门禁、审计字段、敏感数据最小化"],
            ["需求第 10-13 章：页面、非功能、验收、上线兼容", "阶段 4/7", "前端状态、性能安全、验收、迁移和回滚"],
            ["设计第 3-5 章：领域模型、迁移、判定算法", "阶段 1/3", "表结构、公共类型、AuthorizationService、DataScopeFilter"],
            ["设计第 6-8 章：Java 接口、REST、登录会话安全", "阶段 2-4", "公共 API、管理接口、会话、安全响应头"],
            ["设计第 9-12 章：缓存、审计、前端、迁移上线、测试", "阶段 4-7", "缓存版本、审计、前端、灰度、回滚和测试"],
        ],
        [2600, 1700, 5060],
        font_size=8.8,
    )

    add_source_note(doc, "本计划基于需求规格说明书 V1.0 和详细设计说明书 V1.0 编制；排期为研发估算，需在阶段 0 结合实际模块数量、存量数据质量和团队投入复核。")

    # Keep the document metadata neutral and suitable for internal circulation.
    props = doc.core_properties
    props.title = "杰事达材料研发系统_系统权限管理功能开发步骤计划_V1.0"
    props.subject = "系统权限管理功能分阶段开发、测试、灰度和回滚计划"
    props.author = "项目组"
    props.comments = "基于需求规格说明书与详细设计说明书编制"
    props.created = datetime.now()
    props.modified = datetime.now()

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    build_doc()
