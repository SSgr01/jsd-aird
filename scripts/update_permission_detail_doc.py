from copy import deepcopy
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(r"G:\Projects\jsd-aird")
SOURCE = ROOT / "docs" / "杰事达材料研发系统_系统权限管理详细设计说明书_V1.0.docx"
OUT = ROOT / "docs" / "杰事达材料研发系统_系统权限管理详细设计说明书_V1.1.docx"

BLUE = "2E74B5"
NAVY = "0B2545"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "EFF6FF"


def set_run_font(run, size=9.2, color="222222", bold=False, italic=False):
    run.font.name = "Microsoft YaHei"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), "Microsoft YaHei")
    rfonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=120, bottom=70, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    total = sum(widths)

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for node in list(grid):
        grid.remove(node)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_border(cell, color="D0D5DD", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def mark_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def clear_cell(cell):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    return p


def add_table(doc, headers, rows, widths, font_size=8.7, header_fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, widths)
    mark_repeat_header(table.rows[0])
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, header_fill)
        set_cell_border(cell)
        p = clear_cell(cell)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(str(text))
        set_run_font(run, size=font_size, color=NAVY, bold=True)

    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            set_cell_border(cell)
            text, fill = value if isinstance(value, tuple) else (value, None)
            if fill:
                set_cell_shading(cell, fill)
            pieces = str(text).split("\n")
            for part_index, part in enumerate(pieces):
                p = clear_cell(cell) if part_index == 0 else cell.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                run = p.add_run(part)
                set_run_font(run, size=font_size)
    set_table_geometry(table, widths)
    return table


def add_box(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    set_cell_border(cell, color="B9D6F2", size="6")
    p = clear_cell(cell)
    r = p.add_run(label + "  ")
    set_run_font(r, size=9.2, color=BLUE, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=9.2, color=NAVY)
    return table


def add_text(doc, text, style="Normal"):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, size=10.2 if style == "Normal" else 9.2)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, size=10.0)
    return p


def move_before(anchor_paragraph, element):
    anchor_paragraph._p.addprevious(element)


def find_table(doc, marker):
    for table in doc.tables:
        if marker in "\n".join(cell.text for row in table.rows for cell in row.cells):
            return table
    return None


def update_existing_content(doc):
    # Metadata table.
    metadata = doc.tables[0]
    for row in metadata.rows:
        if row.cells[0].text.strip() == "版本":
            row.cells[1].text = "V1.1"
        if row.cells[0].text.strip() == "状态":
            row.cells[1].text = "设计补充"

    # Version history table.
    history = doc.tables[1]
    history.add_row()
    new_row = history.rows[-1]
    values = ["V1.1", "2026-08-19", "明确技术栈、认证会话、持久化、缓存/审计、前端、测试和部署选择，补充替代方案与实现边界。", "设计补充"]
    for cell, value in zip(new_row.cells, values):
        cell.text = value

    # Title/subtitle and static contents summary.
    if len(doc.paragraphs) >= 3:
        doc.paragraphs[2].text = "角色基线、个人覆盖、数据范围过滤与一期技术选型的技术实现设计"
    contents = find_table(doc, "设计概述")
    if contents:
        for row in contents.rows:
            if row.cells[0].text.strip() == "1":
                row.cells[2].text = "设计决策、技术选型、总体目标和技术取舍"

    # Add explicit choices to the original design decision table.
    decisions = find_table(doc, "授权主模型")
    if decisions:
        additions = [
            ("后端形态", "Java 21 + Spring Boot 3.5 + Spring Modulith 模块化单体", "与现有仓库、模块边界和单 JAR 部署一致；一期不引入微服务运维成本。"),
            ("认证会话", "Spring Security 6 + 自建 iam.login_session 的 opaque token 会话", "可复用安全过滤链和 CSRF/安全头能力，同时满足立即撤销、auth_version 和审计要求。"),
            ("持久化", "PostgreSQL + MyBatis-Plus + Flyway", "与现有 POM、Schema 隔离和迁移基线一致；权限判定不引入第二套数据库。"),
            ("权限缓存", "Caffeine 本地缓存 + PostgreSQL policyVersion/authVersion 作为真源", "当前工程不运行 Redis；版本校验保证撤权/禁用即时生效，多实例通过既有事件/版本机制同步。"),
            ("前端状态", "React + TypeScript + Zustand + Axios + Ant Design", "与现有前端依赖一致；状态层只做体验控制，后端 AuthorizationService 仍是最终边界。"),
        ]
        for values in additions:
            row = decisions.add_row()
            for cell, value in zip(row.cells, values):
                cell.text = value

    # Make the existing security requirements concrete.
    security = find_table(doc, "密码存储")
    if security:
        for row in security.rows:
            key = row.cells[0].text.strip()
            if key == "密码存储":
                row.cells[1].text = "使用 Spring Security Argon2PasswordEncoder（Argon2id）；参数由安全基线配置，禁止明文、可逆加密和日志输出。"
            elif key == "会话":
                row.cells[1].text = "使用 Spring Security SecurityFilterChain + 自定义 SessionAuthenticationFilter；iam.login_session 保存不可逆 token_hash，浏览器使用 HttpOnly、Secure、SameSite Cookie。"
            elif key == "CSRF":
                row.cells[1].text = "使用 Spring Security CSRF；前端通过 CookieCsrfTokenRepository 获取 XSRF-TOKEN 并在状态变更请求携带 X-XSRF-TOKEN。"
            elif key == "开发身份":
                row.cells[1].text = "DevelopmentIdentityFilter 仅 local profile 装配；staging/prod 的 SecurityFilterChain 拒绝 X-User-Id、X-Organization-Id、X-Username。"


def insert_technical_selection(doc):
    anchor = next((p for p in doc.paragraphs if p.text.strip() == "2. 系统架构与模块边界"), None)
    if anchor is None:
        raise RuntimeError("未找到章节 2 标题，无法插入技术选型章节")

    blocks = []
    blocks.append(add_heading(doc, "1.2. 技术选型与实现边界", 1))
    blocks.append(add_text(doc, "本节将本期权限管理的技术栈、运行形态、关键依赖和不选方案固化为实现基线。后续如需替换技术，必须同步更新 pom.xml/package.json、ADR、数据库迁移、接口契约和测试门禁；仅修改设计文字不视为完成技术变更。"))
    blocks.append(add_box(doc, "选型结论", "一期沿用现有仓库技术栈，不新增微服务、独立 IAM 服务、Redis、消息队列或 Kubernetes。权限核心作为 iam 模块落在 Spring Boot 模块化单体中，采用 Spring Security 安全过滤链 + PostgreSQL 中的自建 opaque session + MyBatis-Plus/Flyway 持久化；缓存采用 Caffeine 本地缓存，数据库权限版本是最终真源。"))
    blocks.append(add_heading(doc, "1.2.1. 技术选型决策矩阵", 2))
    matrix_rows = [
        ["后端运行时", "Java 21；Spring Boot 3.5.16；Spring MVC、Validation、Actuator", "现有 jsd-aird-api/pom.xml 已锁定；统一依赖、探活和配置模型。", "不选微服务运行时、Quarkus、Micronaut；仍输出单个 Spring Boot JAR。"],
        ["模块边界", "Spring Modulith 1.4.12 + 单仓库模块化单体", "与 ADR-0001 和显式模块检测策略一致，保留模块间 API 边界。", "不把 iam 拆成独立服务；跨模块只依赖 iam::api / ops::api。"],
        ["HTTP/API", "Spring MVC REST + Bean Validation + Springdoc OpenAPI 2.8.17", "与现有 Controller/Swagger 基线一致，适合管理端和业务模块接入。", "不引入 GraphQL/BFF；API 版本固定为 /api/v1。"],
        ["认证与会话", "Spring Security 6（由 Spring Boot 管理）+ 自定义数据库会话过滤器", "复用安全过滤链、CSRF、安全响应头和会话固定防护；会话要求可立即撤销并可按 auth_version 失效。", "不选 JWT 作为一期浏览器会话，不引入 Spring Session/Redis；token 仅存哈希。"],
        ["密码与登录防护", "Spring Security Argon2PasswordEncoder（Argon2id）+ 账号/IP 窗口限流", "满足强哈希、失败计数、短时锁定和审计要求，避免自研密码算法。", "不选明文、MD5/SHA-1、可逆加密或前端哈希代替服务端哈希。"],
        ["关系数据库", "PostgreSQL 18 为 CI/Compose 基线；本地兼容 PostgreSQL 16-18", "现有工程和 schema 隔离均以 PostgreSQL 为基础，JSONB 可承载审计详情。", "不新增 MySQL、Oracle 或权限专用数据库。"],
        ["持久化访问", "MyBatis-Plus 3.5.17 + Mapper/Repository/Converter/TypeHandler 分层", "与 ADR-0003 和现有持久化规则一致，便于把 organization_id、DataScopeFilter 和审计条件显式写入查询。", "不引入 JPA/Hibernate；应用层和领域层不得依赖 Mapper。"],
        ["数据库迁移", "Flyway + 现有 db/migration 目录", "现有迁移链已在使用，适合新增 IAM 表、索引、种子数据和可重复升级。", "不选 Liquibase；不得用启动时自动建表替代版本化迁移。"],
        ["权限缓存", "Spring Cache 抽象 + Caffeine 本地缓存；policyVersion/authVersion 版本校验", "当前仓库无 Redis；本地缓存可降低判权开销，版本真源保证撤权、禁用和密码重置不被旧缓存阻塞。", "多实例通过数据库版本或 ops Outbox/事件传播失效；达到规模后再评估 Redis。"],
        ["审计与观测", "ops.audit_log（PostgreSQL JSONB）+ SLF4J/Logback + Spring Boot Actuator/Micrometer", "审计需要持久化、可按主体/动作/资源检索；运行指标复用现有 Actuator。", "不把普通应用日志当作唯一审计；一期不新增 Kafka/ELK/Prometheus 部署依赖。"],
        ["AI 外发", "Spring AI 1.1.8 / OpenAI-compatible 客户端前置统一 IAM 与数据范围门禁", "现有 AI 依赖已在 POM；把普通查看和 AI 外发权限分开，避免绕过资源范围。", "不允许业务模块直接调用外部模型 SDK；不允许先外发再补审计。"],
        ["前端", "React 18.3.1 + TypeScript 5.9.3 + Vite 6.4.3 + React Router 6.30.4 + Ant Design 5.29.3 + Zustand 5.0.14 + Axios 1.19.0", "与现有 jsd-aird-web/package.json 一致，减少新依赖和状态管理分裂。", "不引入 Redux 或第二套 UI 组件库；前端权限只用于体验优化。"],
        ["测试", "JUnit 5/Spring Boot Test + Testcontainers PostgreSQL + ArchUnit；Vitest/Testing Library/Playwright", "覆盖算法、真实数据库约束、模块边界、页面状态和端到端越权场景。", "不以 Mock 数据库代替迁移/范围过滤集成测试；安全测试必须直接调用 API。"],
        ["文件与下载", "MinIO Java SDK 8.5.17 经 ops 统一封装；签名 URL 生成前做对象级授权", "复用现有文件能力，避免绕过 IAM 直接暴露存储地址。", "禁止把 MinIO bucket 设置为公共读；下载/导出必须有权限和审计。"],
        ["部署与环境", "单 JAR + PostgreSQL；Windows 本地使用本机 PostgreSQL，Linux/CI 可用 Compose PostgreSQL 18", "符合 README 当前运行基线，降低一期环境变量和运维复杂度。", "不新增 Worker、RabbitMQ、Kubernetes 或 Helm 作为权限功能前置。"],
    ]
    blocks.append(add_table(doc, ["技术域", "本期选型", "选择理由", "替代方案与边界"], matrix_rows, [1200, 2650, 2900, 2610], font_size=7.8))
    blocks.append(add_heading(doc, "1.2.2. 关键实现落点", 2))
    implementation_rows = [
        ["Spring Security + DB 会话", "新增 security filter chain、SessionAuthenticationFilter、PasswordEncoder、CSRF 配置；会话写入 iam.login_session。", "生产只认 HttpOnly Cookie；禁止 X-User-* 请求头；auth_version/revoked_at/过期时间每次校验。"],
        ["MyBatis-Plus + Flyway", "iam/infrastructure/persistence；db/migration 新增 IAM 表、索引、种子数据和版本约束。", "查询必须带 organization_id；应用层/领域层不得引用 Mapper。"],
        ["AuthorizationService", "iam/api 公开 AuthorizationService、PermissionDefinitionRegistry、DataScopeFilter；业务模块仅依赖 iam::api。", "动作授权和数据范围分离；列表/搜索/导出在 Repository 层过滤。"],
        ["Caffeine + 版本真源", "在 iam application/infrastructure 增加缓存适配；键包含 organizationId/userId/policyVersion/authVersion。", "缓存只放权限和范围摘要，不放业务敏感数据；版本变化必须失效/回源。"],
        ["ops 审计 + Actuator", "权限变更、拒绝、高风险允许、登录失败和会话撤销调用 ops::api；Actuator 暴露健康和指标。", "审计记录 before/after 摘要、reasonCode、traceId、path 和策略版本；日志不得写原始敏感材料。"],
        ["React/Zustand/Axios", "AuthStore 通过 /api/v1/auth/me 恢复；PermissionService 提供 can/canAny/canAll；Axios 统一 401/403/409。", "路由、菜单、按钮隐藏不能取代后端检查；保存覆盖时携带 expectedVersion。"],
    ]
    blocks.append(add_table(doc, ["选型", "代码/配置落点", "强制约束"], implementation_rows, [1900, 4300, 3160], font_size=8.2, header_fill=LIGHT_GRAY))
    blocks.append(add_heading(doc, "1.2.3. 本期明确不引入的技术", 2))
    for text in [
        "不引入 JWT 作为浏览器会话：一期要求服务端立即撤销、密码/权限变更失效和审计，opaque token + 数据库会话更直接。",
        "不引入 Redis：当前工程没有 Redis 运行基线；使用 Caffeine + PostgreSQL 版本真源，后续按实例数和判权 QPS 评估升级。",
        "不引入微服务或独立 IAM：现阶段权限与业务模块同仓库、同数据库事务和同一发布单，模块化单体足以支撑一期。",
        "不引入第二套 ORM、消息总线或专用审计平台：避免权限核心出现重复持久化、重复事件和难以回滚的外部依赖。",
    ]:
        blocks.append(add_bullet(doc, text))
    blocks.append(add_heading(doc, "1.2.4. 技术选型变更门禁", 2))
    blocks.append(add_text(doc, "任何将 PostgreSQL、MyBatis-Plus、Spring Security 会话、Caffeine、Spring Modulith、React/Zustand/Axios 或现有测试/部署基线替换为其他技术的变更，必须提交 ADR，说明性能、安全、迁移、回滚和运维影响，并由技术负责人和产品负责人共同批准。"))

    # Move the newly created blocks, preserving their authored order.
    for block in blocks:
        element = block._p if hasattr(block, "_p") else block._tbl
        move_before(anchor, element)


def main():
    doc = Document(str(SOURCE))
    update_existing_content(doc)
    insert_technical_selection(doc)
    doc.core_properties.title = "杰事达材料研发系统_系统权限管理详细设计说明书_V1.1"
    doc.core_properties.subject = "统一权限管理详细设计与明确技术选型"
    doc.core_properties.comments = "V1.1 补充一期技术栈、实现落点、替代方案和变更门禁"
    doc.core_properties.modified = datetime.now()
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
