from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt

from update_permission_detail_doc import (
    LIGHT_GRAY,
    LIGHT_BLUE,
    NAVY,
    add_bullet,
    add_box,
    add_heading,
    add_table,
    add_text,
    clear_cell,
    find_table,
    move_before,
    set_cell_border,
    set_cell_width,
    set_run_font,
    set_table_geometry,
)


ROOT = Path(r"G:\Projects\jsd-aird")
SOURCE = ROOT / "docs" / "杰事达材料研发系统_系统权限管理详细设计说明书_V1.2.docx"
OUT = ROOT / "docs" / "杰事达材料研发系统_系统权限管理详细设计说明书_V1.3.docx"
PROTOTYPE = r"G:\Projects\jsd-aird\docs\杰事达材料研发系统_原型终.html"


def update_metadata(doc):
    metadata = doc.tables[0]
    for row in metadata.rows:
        if row.cells[0].text.strip() == "版本":
            row.cells[1].text = "V1.3"
        elif row.cells[0].text.strip() == "状态":
            row.cells[1].text = "设计补充"

    history = doc.tables[1]
    row = history.add_row()
    for cell, value in zip(row.cells, [
        "V1.3",
        "2026-08-19",
        "补充原型视觉与交互基线；明确用户管理、权限管理、登录页和统一布局的实现约束与验收项。",
        "设计补充",
    ]):
        cell.text = value

    # Keep the document's opening statement aligned with the actual local prototype.
    if len(doc.paragraphs) > 2:
        doc.paragraphs[2].text = "角色基线、个人覆盖、数据范围过滤、Spring Security 与原型一致性设计"
    if len(doc.paragraphs) > 3:
        doc.paragraphs[3].text = "内部使用 · 前端视觉与交互以原型终.html为基线，权限判定以本详细设计和正式接口为准"

    # Make the existing front-end section explicitly point to the prototype baseline.
    for para in doc.paragraphs:
        if para.text.strip() == "10. 前端详细设计":
            para.text = "10. 前端详细设计（按原型视觉与交互基线实现）"
            break

    frontend = find_table(doc, "前端能力")
    if frontend is not None and not any("原型视觉基线" in row.cells[0].text for row in frontend.rows):
        values = [
            "原型视觉基线",
            "复用原型的页面壳层、设计令牌、卡片/表格/弹窗/抽屉/Toast 和响应式规则；Ant Design 只作为组件实现载体，必须通过主题令牌覆盖默认样式。",
        ]
        row = frontend.add_row()
        for cell, value in zip(row.cells, values):
            p = clear_cell(cell)
            run = p.add_run(value)
            set_run_font(run, size=8.2)
            set_cell_border(cell)
        set_table_geometry(frontend, [2200, 7160])

    doc.core_properties.title = "杰事达材料研发系统_系统权限管理详细设计说明书_V1.3"
    doc.core_properties.subject = "统一权限管理详细设计、Spring Security 与原型一致性实现方案"
    doc.core_properties.comments = "V1.3 补充原型视觉、布局和用户/权限管理页面设计"
    doc.core_properties.modified = datetime.now()


def insert_prototype_frontend_design(doc):
    anchor = next((p for p in doc.paragraphs if p.text.strip() == "11. 迁移、上线与回滚"), None)
    if anchor is None:
        raise RuntimeError("未找到第 11 章标题，无法插入原型前端详细设计")

    blocks = []
    blocks.append(add_heading(doc, "10.3. 原型视觉与交互基线", 2))
    blocks.append(add_text(doc, f"本节以本地原型文件 {PROTOTYPE} 作为用户管理、权限管理、登录页和公共布局的视觉与交互基线。原型当前主要展示工作台及业务页面，脚本中已删除独立登录页；因此登录页复用原型的设计令牌和表单/卡片语言实现，不能声称存在可逐像素复制的登录稿。功能、字段、授权结果和安全行为仍以需求规格说明书、本文档及后端接口契约为准。"))
    blocks.append(add_box(doc, "实现原则", "页面可以复用原型的结构、颜色、间距和交互反馈，但前端隐藏菜单、按钮禁用和 localStorage 演示数据都不构成安全边界；所有用户、角色、个人覆盖、数据范围和 AI 外发权限必须由后端 AuthorizationService 最终判定。"))

    blocks.append(add_heading(doc, "10.3.1. 原型设计令牌", 3))
    token_rows = [
        ["页面背景/表面", "#f3f6fb / #ffffff", "工作区使用浅灰蓝背景，卡片、弹窗和抽屉使用白色表面；禁止在权限页面另起一套背景色。"],
        ["主色", "#2f66e8；悬停 #2453c6；浅色 #edf3ff", "主按钮、链接、激活 Tab、选中项和焦点状态统一使用主色体系。"],
        ["文本", "#17233a / #66758b / #9aa7b8", "标题、正文、辅助说明和占位符分别使用主文本、次文本和三级文本。"],
        ["边框/语义色", "#dfe7f2；成功 #16a34a；警告 #f59e0b；危险 #dc2626；紫色 #7c3aed", "状态标签、DENY、高风险 AI 外发和提示信息使用语义色，不用颜色单独表达关键权限结论。"],
        ["侧边栏", "linear-gradient(180deg,#2f63d9 0%,#4c8ddd 62%,#76ccc6 100%)", "公共导航延续原型蓝到青的渐变；激活项使用白色文字/图标和半透明绿色背景。"],
        ["尺寸", "侧栏 224px；收起 64px；顶栏 58px；内容内边距 24px；最大内容宽 1440px", "用户管理、权限管理和登录页在不同屏幕下保持同一比例与对齐规则。"],
        ["圆角/阴影", "6px / 10px / 14px；1px、7px、18px 阴影层级", "输入框/按钮、卡片、弹窗或大容器分别使用对应层级；避免新增厚重阴影。"],
        ["字体/动效", "Inter、Segoe UI、系统无衬线；200ms cubic-bezier(0.4,0,0.2,1)", "中文使用系统回退字体；hover、展开、Toast 和抽屉过渡统一 200ms。"],
    ]
    blocks.append(add_table(doc, ["设计项", "原型基线", "实现约束"], token_rows, [1800, 3000, 4560], font_size=7.8, header_fill=LIGHT_BLUE))

    blocks.append(add_heading(doc, "10.3.2. 公共页面壳层与布局", 3))
    for text in [
        "应用采用 AppLayout：左侧 Sidebar + 顶部 Topbar + 主内容 PageContainer。Sidebar 默认 224px，可收起为 64px；Topbar 高度 58px；主内容背景为 #f3f6fb，页面内容最大宽度 1440px、内边距 24px。",
        "Topbar 保留原型中的面包屑、通知、快捷入口、头像和用户菜单位置；用户管理和权限管理页面不得移除面包屑，否则返回路径和上下文会丢失。",
        "页面主体统一使用卡片承载功能区：卡片标题区 16px/20px，卡片内容区 20px，底部操作区 12px/20px。筛选区、数据表、差异对比区和权限范围编辑区按卡片分组。",
        "侧栏只负责导航和体验层菜单展示；路由守卫只做页面级体验优化，API 仍必须经过认证和 AuthorizationService。无权限菜单隐藏时，直接访问地址仍应收到后端 401/403。",
        "在 900px 以下优先收起侧栏；在 640px 以下将筛选项换行、表格操作收纳到更多菜单，保留主操作和危险操作的可发现性；不得通过横向溢出隐藏权限结果。",
    ]:
        blocks.append(add_bullet(doc, text))

    blocks.append(add_heading(doc, "10.3.3. 用户管理页面", 3))
    blocks.append(add_text(doc, "用户管理页面沿用原型的页面标题、工具栏、卡片和数据表语言，不单独设计后台管理风格。列表页负责查询和状态操作，新增/编辑用户使用弹窗，复杂的会话和权限摘要使用右侧抽屉；所有写操作都要对接 IAM API 并显示明确反馈。"))
    user_rows = [
        ["页面区块", "布局与组件", "字段/行为约束"],
        ["页面头部", "PageContainer + 标题/说明 + 右侧主按钮", "主按钮使用 accent；新增用户前端可按权限隐藏，但后端必须再次校验 USER_CREATE。"],
        ["筛选工具栏", "搜索框、组织/角色选择、状态 Tabs 或 Select、查询/重置按钮", "搜索 username/displayName；查询条件不能突破组织隔离；筛选条件改变时清晰展示当前条件。"],
        ["用户表格", "复用 data-table：用户、部门、主角色、状态、最后登录、更新时间、操作", "状态使用语义 Badge；最后登录为空显示“从未登录”；操作列不得只依赖图标，危险操作需要文字或 Tooltip。"],
        ["新增/编辑", "复用 modal-box 视觉，最大宽度 600px；表单使用 form-input 规格", "用户名不可修改或按后端规则处理；密码只在创建/重置流程出现；组织、主角色和状态必走服务端校验。"],
        ["会话与重置", "复用 420px drawer 展示登录会话、最后活动和撤销操作；确认动作使用 danger modal", "重置密码、禁用账号、强制下线均需二次确认；成功后刷新用户摘要并提示会话已失效。"],
        ["反馈状态", "loading skeleton、空态、无权限态、错误态、Toast", "空数据与无权限必须区分；401 进入登录页，403 显示无权原因，409 显示并发冲突及重新加载入口。"],
    ]
    blocks.append(add_table(doc, ["页面区块", "布局与组件", "字段/行为约束"], user_rows[1:], [1600, 3200, 4560], font_size=7.7, header_fill=LIGHT_BLUE))
    blocks.append(add_box(doc, "用户页安全约束", "停用、密码重置、强制下线和主角色变更均属于高影响操作；页面上必须显示影响范围和二次确认，后端在事务内更新状态/版本并撤销旧会话，不能只刷新前端表格。"))

    blocks.append(add_heading(doc, "10.3.4. 权限管理页面", 3))
    blocks.append(add_text(doc, "权限管理页面使用“筛选上下文 + 权限卡片/矩阵 + 编辑抽屉”的结构。角色默认、个人覆盖和生效差异是三个可切换视图；动作授权与数据范围保持同一行的上下文关系，避免把 ALL/SELF/ASSIGNED 等范围误解为独立权限。"))
    permission_rows = [
        ["视图", "原型化布局", "关键交互与约束"],
        ["角色默认权限", "Tabs + 角色选择器 + 权限矩阵卡片", "按 permissionCode 分组展示动作；显示默认 ALLOW/DENY、DataScope 和继承来源；保存必须携带 expectedVersion。"],
        ["个人权限覆盖", "用户选择器 + 差异列表/矩阵 + 编辑 Drawer", "每条覆盖显示 inherited/override 来源；ALLOW/DENY 和范围编辑放在同一上下文；恢复角色默认必须可撤销。"],
        ["生效权限差异", "左右或上下差异卡片，使用浅色 accent 和语义标签", "明确展示角色基线、个人覆盖、生效结果和冲突来源；不能只展示最终布尔值。"],
        ["数据范围编辑", "Drawer 或 modal 中使用 Select、Tree/TreeSelect、已选目标 Chips", "ALL/SELF/ASSIGNED/PROJECT/CATEGORY/SELECTED 使用枚举选项；SELECTED 无目标时保存按钮禁用并提示原因。"],
        ["高风险动作", "危险按钮、警告提示、敏感外发标记", "DENY、高风险动作和 AI 外发权限使用 danger/warning 语义色；颜色只是辅助，文字必须明确说明影响。"],
        ["并发与审计", "保存冲突 Toast + 差异 Drawer；变更记录 Tab 或链接", "POLICY_VERSION_CONFLICT 必须保留用户编辑内容的可恢复入口；页面提供最近变更人、时间和 reasonCode。"],
    ]
    blocks.append(add_table(doc, ["视图", "原型化布局", "关键交互与约束"], permission_rows[1:], [1600, 3200, 4560], font_size=7.7, header_fill=LIGHT_BLUE))
    blocks.append(add_box(doc, "权限页边界", "Ant Design Table、Modal、Drawer、Tree 和 Form 仅负责呈现与收集输入；不在前端拼接最终权限、不把 localStorage 里的演示权限当作有效权限、不通过按钮隐藏替代后端校验。"))

    blocks.append(add_heading(doc, "10.3.5. 登录页样式与会话交互", 3))
    blocks.append(add_text(doc, "原型脚本中的 showLoginModal() 明确表示已删除登录页并直接进入工作台，因此本期没有可逐像素复制的登录页画面。实现时新建独立 LoginPage，但只复用原型已确定的设计语言：#f3f6fb 页面背景、白色卡片、14px 大圆角、浅阴影、#2f66e8 主按钮、6px 输入框圆角、Inter/Segoe UI 字体和 200ms 过渡。"))
    login_rows = [
        ["布局", "全屏居中单列卡片，建议宽度 420px；登录页不显示业务侧栏和业务 Topbar", "保持认证场景聚焦；卡片可包含产品标识、标题、账号、密码、登录按钮和错误提示。"],
        ["表单", "复用 form-input：高度、边框、圆角、焦点主色与原型一致", "username/password 必填；密码不回显；浏览器自动填充不应破坏样式；支持 Enter 提交。"],
        ["状态", "按钮 loading、字段错误、全局错误 Toast/Inline Alert", "AUTH_INVALID 不区分用户不存在和密码错误；账号停用、锁定、CSRF 失败分别显示可行动提示。"],
        ["会话", "成功后进入原型同款 AppLayout，并恢复面包屑/侧栏/用户菜单", "JSD_SESSION 只由服务端通过 HttpOnly Cookie 设置；前端禁止把 token 写入 localStorage/sessionStorage。"],
        ["安全", "不在页面展示 token、auth_version 或密码策略细节", "登录、退出、锁定、强制下线由 Spring Security + 数据库会话承接；登录页不实现 SSO/OAuth2/外部身份源入口。"],
    ]
    blocks.append(add_table(doc, ["设计项", "原型化表达", "实现约束"], login_rows, [1500, 3600, 4260], font_size=7.9, header_fill=LIGHT_BLUE))

    blocks.append(add_heading(doc, "10.3.6. 原型组件到 React 实现映射", 3))
    mapping_rows = [
        ["原型模式", "React 实现", "实现要求"],
        ["#app / Sidebar / Topbar", "AppLayout、Sidebar、Topbar、PageContainer", "统一管理收起状态、面包屑、当前用户和响应式断点；不在业务页复制壳层。"],
        [".card / .card-header / .card-body", "Ant Design Card 或自定义 PageCard", "通过 CSS variables/ConfigProvider 覆盖颜色、圆角、边框、内边距和阴影，保持原型密度。"],
        [".data-table", "Ant Design Table", "统一表头背景、行高、hover、分页和操作列；表格必须支持 loading/empty/error/forbidden。"],
        [".modal-box", "Ant Design Modal", "最大宽度按 600px 约束；危险操作使用确认文案和 danger 按钮。"],
        [".drawer", "Ant Design Drawer", "权限差异、会话详情和数据范围编辑优先使用右侧 420px 抽屉，保存/取消固定在底部操作区。"],
        [".form-input / select", "Ant Design Form、Input、Password、Select", "统一 label、help、error、focus 状态；校验错误需能定位到字段。"],
        [".tree-nav / tree-item", "Ant Design Tree/TreeSelect 或封装 TreeFilter", "组织、项目、分类和 SELECTED 目标共用树选择约定，保留选中、禁用和半选状态。"],
        [".toast-item", "AppNotification / message / notification 封装", "统一成功、警告、错误、无权限和冲突反馈；不让各页面直接拼接不同提示风格。"],
    ]
    blocks.append(add_table(doc, ["原型模式", "React 实现", "实现要求"], mapping_rows[1:], [2200, 2600, 4560], font_size=7.7, header_fill=LIGHT_BLUE))
    blocks.append(add_text(doc, "前端当前技术基线为 React 18 + TypeScript + Vite + Ant Design + Zustand + Axios；因此本节不改造为 Vue，也不引入另一套 UI 框架。建议在 jsd-aird-web/src/theme/prototype-tokens.ts 集中声明颜色、圆角、尺寸和阴影，在 AppLayout 及 IAM 页面统一引用。"))

    blocks.append(add_heading(doc, "10.3.7. 原型对照验收清单", 3))
    acceptance_rows = [
        ["UX-001", "公共壳层", "侧栏 224/64px 收起、顶栏 58px、面包屑、页面背景、内容宽度和 24px 内边距与原型一致。"],
        ["UX-002", "设计令牌", "主色、文本色、边框色、语义色、圆角、阴影、字体和 200ms 过渡统一从原型令牌生成。"],
        ["UX-003", "用户列表", "用户管理列表沿用原型卡片/工具栏/表格密度；加载、空态、无权、异常和冲突状态可区分。"],
        ["UX-004", "用户编辑", "新增/编辑使用 600px 级别弹窗，会话/复杂信息使用 420px 抽屉；危险动作有确认和结果反馈。"],
        ["UX-005", "权限矩阵", "角色默认、个人覆盖、生效差异有清晰 Tab/卡片层次；ALLOW/DENY、继承来源、DataScope 同时可见。"],
        ["UX-006", "范围选择", "ALL/SELF/ASSIGNED/PROJECT/CATEGORY/SELECTED 使用统一选择组件；SELECTED 为空时阻止保存并说明原因。"],
        ["UX-007", "登录页", "登录页复用原型设计系统；不出现 SSO/OAuth2 入口，不把 token 放进浏览器存储，不显示业务侧栏。"],
        ["UX-008", "响应式", "900px 以下侧栏收起，640px 以下筛选换行且关键操作可见；权限结论不能被横向溢出遮挡。"],
        ["SEC-UX-009", "后端边界", "隐藏菜单、禁用按钮和前端 can() 只优化体验；直接调用受保护 API 仍由 Spring Security/IAM 返回准确 401/403。"],
    ]
    blocks.append(add_table(doc, ["编号", "验收范围", "验收标准"], acceptance_rows, [1500, 2200, 5660], font_size=7.8, header_fill=LIGHT_GRAY))

    for block in blocks:
        element = block._p if hasattr(block, "_p") else block._tbl
        move_before(anchor, element)


def main():
    doc = Document(str(SOURCE))
    update_metadata(doc)
    insert_prototype_frontend_design(doc)
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
